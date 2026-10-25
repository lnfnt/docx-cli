"""Tests for document builder and generated file properties."""
from __future__ import annotations

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_cli.builder import build_document
from docx_cli.spec import DocSpec


class TestMarginsAndPage:
    def test_left_margin(self, built_path):
        doc = Document(str(built_path))
        assert abs(doc.sections[0].left_margin.cm - 3.0) < 0.05

    def test_right_margin(self, built_path):
        doc = Document(str(built_path))
        assert abs(doc.sections[0].right_margin.cm - 1.5) < 0.05

    def test_top_margin(self, built_path):
        doc = Document(str(built_path))
        assert abs(doc.sections[0].top_margin.cm - 2.0) < 0.05

    def test_bottom_margin(self, built_path):
        doc = Document(str(built_path))
        assert abs(doc.sections[0].bottom_margin.cm - 2.0) < 0.05

    def test_page_width_a4(self, built_path):
        doc = Document(str(built_path))
        assert abs(doc.sections[0].page_width.cm - 21.0) < 0.05

    def test_page_height_a4(self, built_path):
        doc = Document(str(built_path))
        assert abs(doc.sections[0].page_height.cm - 29.7) < 0.05

    def test_compact_profile_margins(self, tmp_workspace, minimal_spec):
        from docx_cli.spec import Profile

        compact = DocSpec.from_dict(
            {
                **minimal_spec.model_dump(),
                "profile": "compact",
            }
        )
        out = tmp_workspace / "compact.docx"
        build_document(compact, str(out))
        doc = Document(str(out))
        profile = Profile("compact")
        assert abs(doc.sections[0].left_margin.cm - profile.margin_left_cm) < 0.05
        assert abs(doc.sections[0].right_margin.cm - profile.margin_right_cm) < 0.05


class TestHeaderAndTitle:
    def test_header_lines_present(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        texts = [p.text.strip() for p in doc.paragraphs]
        for line in minimal_spec.header.lines:
            assert line in texts

    def test_header_alignment(self, built_path):
        doc = Document(str(built_path))
        header_text = "В тестовый суд"
        for p in doc.paragraphs:
            if p.text.strip() == header_text:
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    def test_title_present(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        assert any(p.text.strip() == minimal_spec.title.main for p in doc.paragraphs)

    def test_title_centered(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            if p.text.strip() == minimal_spec.title.main:
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_title_bold_14pt(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            if p.text.strip() == minimal_spec.title.main:
                for run in p.runs:
                    if run.text:
                        assert run.bold is True
                        assert run.font.size.pt == 14

    def test_title_line_spacing_from_profile(self, built_path, minimal_spec):
        from docx_cli.spec import Profile

        doc = Document(str(built_path))
        profile = Profile(minimal_spec.profile)
        for p in doc.paragraphs:
            if p.text.strip() == minimal_spec.title.main:
                assert abs(float(p.paragraph_format.line_spacing) - profile.header_line_spacing) < 0.05

    def test_subtitle_present(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        assert any(
            p.text.strip() == minimal_spec.title.subtitle for p in doc.paragraphs
        )

    def test_subtitle_centered(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            if p.text.strip() == minimal_spec.title.subtitle:
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER


class TestBodyParagraphs:
    def test_paragraph_text_present(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        assert any(
            p.text.strip() == minimal_spec.body[0].text for p in doc.paragraphs
        )

    def test_paragraph_justified(self, built_path):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_paragraph_line_spacing(self, built_path):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                assert abs(float(p.paragraph_format.line_spacing) - 1.5) < 0.05

    def test_paragraph_first_line_indent(self, built_path):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                assert abs(p.paragraph_format.first_line_indent.cm - 1.25) < 0.05

    def test_body_font_size_12pt(self, built_path):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                for run in p.runs:
                    if run.text:
                        assert run.font.size.pt == 12

    def test_inline_bold(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "**жирный** текст", "inline_markdown": True}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "bold.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if "жирный" in p.text:
                bold_runs = [r for r in p.runs if r.bold and r.text.strip()]
                assert any("жирный" in r.text for r in bold_runs)

    def test_inline_italic(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "*курсив* текст", "inline_markdown": True}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "italic.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if "курсив" in p.text:
                italic_runs = [r for r in p.runs if r.italic and r.text.strip()]
                assert any("курсив" in r.text for r in italic_runs)

    def test_inline_markdown_disabled(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "**буквы**", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "nomd.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if "**" in p.text:
                assert all(not r.bold for r in p.runs if r.text)


class TestNumberedList:
    def test_numbered_items_present(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "numbered_list", "items": ["первый", "второй", "третий"]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "list.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        texts = [p.text.strip() for p in doc.paragraphs]
        assert "1. первый" in texts
        assert "2. второй" in texts
        assert "3. третий" in texts

    def test_numbered_items_justified(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "numbered_list", "items": ["первый"]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "list.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip().startswith("1."):
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_numbered_items_hanging_indent(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "numbered_list", "items": ["первый"]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "list.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip().startswith("1."):
                assert abs(p.paragraph_format.left_indent.cm - 1.25) < 0.05
                assert abs(p.paragraph_format.first_line_indent.cm - (-0.75)) < 0.05


class TestTable:
    def test_table_headers_and_rows(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {
                "type": "table",
                "headers": ["А", "Б"],
                "rows": [["1", "2"], ["3", "4"]],
            }
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert table.rows[0].cells[0].text == "А"
        assert table.rows[1].cells[1].text == "2"

    def test_table_header_bold(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {
                "type": "table",
                "headers": ["А"],
                "rows": [["1"]],
            }
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        header_cell = doc.tables[0].rows[0].cells[0]
        assert any(r.bold for r in header_cell.paragraphs[0].runs if r.text)

    def test_table_widths_cm(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {
                "type": "table",
                "headers": ["А", "Б"],
                "rows": [["1", "2"]],
                "widths_cm": [5.0, 7.0],
            }
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        row = doc.tables[0].rows[0]
        assert abs(row.cells[0].width.cm - 5.0) < 0.1
        assert abs(row.cells[1].width.cm - 7.0) < 0.1

    def test_table_cell_font_size_11pt(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "table", "headers": ["А"], "rows": [["1"]]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table_font.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        cell = doc.tables[0].rows[1].cells[0]
        for run in cell.paragraphs[0].runs:
            if run.text:
                assert run.font.size.pt == 11

    def test_table_cell_line_spacing_115(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "table", "headers": ["А"], "rows": [["1"]]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table_ls.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        cell = doc.tables[0].rows[1].cells[0]
        assert abs(float(cell.paragraphs[0].paragraph_format.line_spacing) - 1.15) < 0.05


class TestSignatureAndAttachments:
    def test_signature_block(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "signature_block", "left": "Исполнитель", "right": "___/Иванов/"}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "sig.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        assert len(doc.tables) == 1
        left = doc.tables[0].rows[0].cells[0].text.strip()
        right = doc.tables[0].rows[0].cells[1].text.strip()
        assert "Исполнитель" in left
        assert "Иванов" in right

    def test_attachments_list(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "attachments_list", "items": ["1. Копия", "2. Выписка"]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "attach.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        texts = [p.text.strip() for p in doc.paragraphs]
        assert "Приложения:" in texts
        assert "1. Копия" in texts
        assert "2. Выписка" in texts

    def test_attachments_items_hanging_indent(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "attachments_list", "items": ["1. Копия"]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "attach.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip() == "1. Копия":
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
                assert abs(p.paragraph_format.left_indent.cm - 1.25) < 0.05
                assert abs(p.paragraph_format.first_line_indent.cm - (-0.75)) < 0.05


class TestFontsGlobal:
    def test_all_non_empty_runs_times_new_roman(self, built_path):
        doc = Document(str(built_path))
        for p in doc.paragraphs:
            for run in p.runs:
                if run.text:
                    assert run.font.name == "Times New Roman"

    def test_core_properties_author(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        assert doc.core_properties.author == minimal_spec.output.author

    def test_core_properties_title(self, built_path, minimal_spec):
        doc = Document(str(built_path))
        assert doc.core_properties.title == minimal_spec.title.main


class TestExample:
    def test_example_builds_and_validates(self, example_spec_path, tmp_workspace):
        import json

        with open(example_spec_path, "r", encoding="utf-8") as f:
            spec = DocSpec.from_dict(json.load(f))
        out = tmp_workspace / "example.docx"
        result = build_document(spec, str(out))
        assert result["ok"] is True
        assert result["status"] == "ok"
        assert out.exists()
        doc = Document(str(out))
        assert any(p.text.strip() == spec.title.main for p in doc.paragraphs)
