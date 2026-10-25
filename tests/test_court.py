"""Tests for the 'court' profile and related court-specific rendering."""
from __future__ import annotations

import json
import zipfile
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_cli.builder import build_document
from docx_cli.errors import SpecError
from docx_cli.spec import DocSpec, Profile


PYTHON = sys.executable


@pytest.fixture
def court_spec_dict():
    return {
        "profile": "court",
        "output": {"file_name": "court.docx", "author": "Адвокат Т.Т."},
        "header": {
            "align": "right",
            "lines": [
                "В тестовый суд",
                "",
                "Истец: Тестов Т.Т.",
                "Адрес: г. Москва",
                "",
                "Ответчик: Иванов И.И.",
            ],
        },
        "title": {"main": "ОТЗЫВ", "subtitle": "на возражения"},
        "body": [
            {"type": "paragraph", "text": "Текст документа для проверки.", "inline_markdown": False}
        ],
        "checks": {"no_yo": True, "no_en_dash": True, "require_norms_refs": False},
    }


@pytest.fixture
def court_spec(court_spec_dict):
    return DocSpec.from_dict(court_spec_dict)


@pytest.fixture
def court_path(tmp_workspace, court_spec):
    out = tmp_workspace / "court.docx"
    build_document(court_spec, str(out))
    return out


class TestCourtProfile:
    def test_court_profile_parses(self, court_spec_dict):
        spec = DocSpec.from_dict(court_spec_dict)
        assert spec.profile == "court"

    def test_court_profile_defaults(self):
        profile = Profile("court")
        assert profile.body_font_size_pt == 14
        assert profile.title_font_size_pt == 14
        assert profile.cell_font_size_pt == 14
        assert profile.body_line_spacing == 1.0
        assert profile.header_line_spacing == 1.0
        assert profile.default_signature_format == "text"

    def test_court_margins_same_as_gost(self, court_path):
        doc = Document(str(court_path))
        section = doc.sections[0]
        assert abs(section.left_margin.cm - 3.0) < 0.05
        assert abs(section.right_margin.cm - 1.5) < 0.05
        assert abs(section.top_margin.cm - 2.0) < 0.05
        assert abs(section.bottom_margin.cm - 2.0) < 0.05

    def test_court_line_spacing_option_115(self, tmp_workspace, court_spec_dict):
        court_spec_dict["line_spacing"] = 1.15
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_115.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                assert abs(float(p.paragraph_format.line_spacing) - 1.15) < 0.05


class TestCourtHeader:
    def test_court_header_left_indent_and_alignment(self, court_path, court_spec):
        doc = Document(str(court_path))
        header_texts = {line.strip() for line in court_spec.header.lines if line != ""}
        found = 0
        for p in doc.paragraphs:
            if p.text.strip() in header_texts:
                found += 1
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
                assert abs(p.paragraph_format.left_indent.cm - 7.0) < 0.05
        assert found == len(header_texts)

    def test_court_header_font_14pt(self, court_path, court_spec):
        doc = Document(str(court_path))
        header_texts = {line.strip() for line in court_spec.header.lines if line != ""}
        for p in doc.paragraphs:
            if p.text.strip() in header_texts:
                for run in p.runs:
                    if run.text:
                        assert run.font.size.pt == 14

    def test_court_header_line_spacing_10(self, court_path, court_spec):
        doc = Document(str(court_path))
        header_texts = {line.strip() for line in court_spec.header.lines if line != ""}
        for p in doc.paragraphs:
            if p.text.strip() in header_texts:
                assert abs(float(p.paragraph_format.line_spacing) - 1.0) < 0.05

    def test_court_header_section_break_space_after(self, court_path):
        doc = Document(str(court_path))
        prev = None
        for p in doc.paragraphs:
            text = p.text.strip()
            if text == "В тестовый суд":
                prev = p
                break
        assert prev is not None
        assert abs(prev.paragraph_format.space_after.pt - 3.0) < 0.1

    def test_court_header_no_empty_paragraph_for_blank_line(self, court_path, court_spec):
        doc = Document(str(court_path))
        # Empty header lines must not produce empty paragraphs.
        empty_count = sum(1 for p in doc.paragraphs if p.text.strip() == "")
        assert empty_count == 0

    def test_court_header_custom_left_indent(self, tmp_workspace, court_spec_dict):
        court_spec_dict["header"]["left_indent_cm"] = 6.5
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_header.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip() == "В тестовый суд":
                assert abs(p.paragraph_format.left_indent.cm - 6.5) < 0.05


class TestCourtTitle:
    def test_court_title_bold_14pt_centered(self, court_path, court_spec):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            if p.text.strip() == court_spec.title.main:
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    if run.text:
                        assert run.bold is True
                        assert run.font.size.pt == 14

    def test_court_title_space_before_after_12pt(self, court_path, court_spec):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            if p.text.strip() == court_spec.title.main:
                assert abs(p.paragraph_format.space_before.pt - 12.0) < 0.1
                assert abs(p.paragraph_format.space_after.pt - 12.0) < 0.1

    def test_court_subtitle_regular_14pt_centered(self, court_path, court_spec):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            if court_spec.title.subtitle and p.text.strip() == court_spec.title.subtitle:
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    if run.text:
                        assert run.bold is False
                        assert run.font.size.pt == 14


class TestCourtBody:
    def test_court_body_font_14pt(self, court_path):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                for run in p.runs:
                    if run.text:
                        assert run.font.size.pt == 14

    def test_court_body_line_spacing_10(self, court_path):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                assert abs(float(p.paragraph_format.line_spacing) - 1.0) < 0.05

    def test_court_body_alignment_justify(self, court_path):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_court_body_first_line_indent_125(self, court_path):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            if p.text.strip() == "Текст документа для проверки.":
                assert abs(p.paragraph_format.first_line_indent.cm - 1.25) < 0.05

    def test_court_all_runs_times_new_roman(self, court_path):
        doc = Document(str(court_path))
        for p in doc.paragraphs:
            for run in p.runs:
                if run.text:
                    assert run.font.name == "Times New Roman"


class TestCourtLists:
    def test_court_numbered_list_hanging_indent(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {"type": "numbered_list", "items": ["первый пункт", "второй пункт"]}
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_list.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip().startswith(("1.", "2.")):
                assert abs(p.paragraph_format.left_indent.cm - 1.25) < 0.05
                assert abs(p.paragraph_format.first_line_indent.cm - (-0.75)) < 0.05

    def test_court_numbered_list_multiline_hanging(self, tmp_workspace, court_spec_dict):
        long_text = "а " * 300
        court_spec_dict["body"] = [{"type": "numbered_list", "items": [long_text]}]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_long_list.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip().startswith("1."):
                assert abs(p.paragraph_format.left_indent.cm - 1.25) < 0.05
                assert abs(p.paragraph_format.first_line_indent.cm - (-0.75)) < 0.05

    def test_court_attachments_list_hanging_indent(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {"type": "attachments_list", "items": ["1. Копия", "2. Выписка"]}
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_attach.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            text = p.text.strip()
            if text.startswith(("1.", "2.")):
                assert abs(p.paragraph_format.left_indent.cm - 1.25) < 0.05
                assert abs(p.paragraph_format.first_line_indent.cm - (-0.75)) < 0.05


class TestCourtPleadingMarker:
    def test_court_pleading_marker_block_centered_bold_no_indent(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {"type": "pleading_marker", "text": "ПРОШУ СУД:"}
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_pleading.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.strip() == "ПРОШУ СУД:":
                assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
                assert abs(p.paragraph_format.first_line_indent.cm - 0.0) < 0.05
                assert abs(p.paragraph_format.space_before.pt - 12.0) < 0.1
                assert abs(p.paragraph_format.space_after.pt - 6.0) < 0.1
                for run in p.runs:
                    if run.text:
                        assert run.bold is True
                        assert run.font.size.pt == 14

    def test_court_pleading_marker_warning_on_body_paragraph(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {"type": "paragraph", "text": "Прошу суд взыскать", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_pleading_warn.docx"
        result = build_document(spec, str(out))
        assert result["ok"] is True
        assert result["status"] == "degraded"
        assert any("pleading_marker" in w for w in result["metadata"]["warnings"])


class TestCourtSignature:
    def test_court_signature_default_text_format_no_table(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {"type": "signature_block", "left": "Представитель", "right": "___/Иванов/"}
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_sig_text.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        assert len(doc.tables) == 0

    def test_court_signature_text_has_tab_stop_right(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {"type": "signature_block", "left": "Представитель", "right": "___/Иванов/"}
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_sig_text.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        sig_p = None
        for p in doc.paragraphs:
            if "Представитель" in p.text and "___/Иванов/" in p.text:
                sig_p = p
                break
        assert sig_p is not None
        stops = [ts.position.cm for ts in sig_p.paragraph_format.tab_stops]
        assert any(abs(pos - 16.5) < 0.1 for pos in stops)

    def test_court_signature_explicit_table_no_borders(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {
                "type": "signature_block",
                "left": "Представитель",
                "right": "___/Иванов/",
                "format": "table",
            }
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_sig_table.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        assert len(doc.tables) == 1
        # All table cell borders must be w:val="none"
        with zipfile.ZipFile(out) as zf:
            xml = zf.read("word/document.xml").decode("utf-8")
        for cell_xml in _cell_xmls(xml):
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                m = __import__("re").search(rf'<w:{edge}[^>]*w:val="([^"]+)"', cell_xml)
                if m:
                    assert m.group(1) == "none"


class TestCourtTable:
    def test_court_table_cell_font_14pt(self, tmp_workspace, court_spec_dict):
        court_spec_dict["body"] = [
            {"type": "table", "headers": ["Статья"], "rows": [["309 ГК РФ"]]}
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_table.docx"
        build_document(spec, str(out))
        doc = Document(str(out))
        cell = doc.tables[0].rows[1].cells[0]
        for run in cell.paragraphs[0].runs:
            if run.text:
                assert run.font.size.pt == 14


class TestCourtValidation:
    def test_court_validator_passes(self, court_path, court_spec):
        from docx_cli.validate import validate_document

        result = validate_document(str(court_path), court_spec)
        assert not result["errors"]
        assert result["checks"]["header_formatting"] == "pass"
        assert result["checks"]["list_formatting"] == "pass"

    def test_court_validator_header_geometry(self, tmp_workspace, court_spec_dict):
        from docx_cli.validate import validate_document

        court_spec_dict["header"]["left_indent_cm"] = 7.5
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_val_header.docx"
        build_document(spec, str(out))
        result = validate_document(str(out), spec)
        assert not any("header_left_indent" in e for e in result["errors"])

    def test_court_validator_lists_pass(self, tmp_workspace, court_spec_dict):
        from docx_cli.validate import validate_document

        court_spec_dict["body"] = [
            {"type": "numbered_list", "items": ["первый", "второй"]},
            {"type": "attachments_list", "items": ["1. Копия"]},
        ]
        spec = DocSpec.from_dict(court_spec_dict)
        out = tmp_workspace / "court_val_lists.docx"
        build_document(spec, str(out))
        result = validate_document(str(out), spec)
        assert result["checks"]["list_formatting"] == "pass"


class TestCourtSpecErrors:
    def test_court_spec_rejects_left_indent_below_range(self, court_spec_dict):
        court_spec_dict["header"]["left_indent_cm"] = 5.9
        with pytest.raises(SpecError):
            DocSpec.from_dict(court_spec_dict)

    def test_court_spec_rejects_left_indent_above_range(self, court_spec_dict):
        court_spec_dict["header"]["left_indent_cm"] = 8.1
        with pytest.raises(SpecError):
            DocSpec.from_dict(court_spec_dict)

    def test_court_spec_rejects_invalid_line_spacing(self, court_spec_dict):
        court_spec_dict["line_spacing"] = 1.2
        with pytest.raises(SpecError):
            DocSpec.from_dict(court_spec_dict)


class TestCourtCli:
    def test_court_profile_cli_override(self, court_spec_dict, tmp_workspace):
        spec_path = tmp_workspace / "court_spec.json"
        spec_path.write_text(json.dumps(court_spec_dict), encoding="utf-8")
        out = tmp_workspace / "court_cli.docx"
        result = __import__("subprocess").run(
            [str(PYTHON), "-m", "docx_cli", "build", str(spec_path), str(out)],
            capture_output=True,
            text=True,
            cwd=str(Path(__file__).parent.parent),
        )
        payload = json.loads(result.stdout.splitlines()[0])
        assert result.returncode == 0
        assert payload["data"]["profile"] == "court"
        assert out.exists()


def _cell_xmls(document_xml: str):
    """Yield each <w:tc> block from the document XML."""
    return __import__("re").findall(r"<w:tc>.*?</w:tc>", document_xml, __import__("re").DOTALL)
