"""docx-cli v0.4.0 приёмка: md-build + normalize (+ scan уже v0.3).

Негативки обязательны: md без титула — fail-closed; normalize меняет шрифты;
иск-каркас собирается с обновлениями по замечаниям юриста (333.21 НК, ст.61-62).
"""
import json
import re
import zipfile
from pathlib import Path

from docx import Document

from docx_cli.md_court import MdBuildError, md_build

ROOT = Path(__file__).resolve().parents[1]


def test_md_build_happy(tmp_path):
    md = "\n".join([
        "В Арбитражный суд города Москвы",
        "Истец: ООО Тест, ИНН 7701001001, ОГРН 1157700000001",
        "Ответчик: ООО Ответчик, ИНН 7702002002, ОГРН 1167700000002",
        "Дело № А40-999/2026",
        "",
        "ИСКОВОЕ ЗАЯВЛЕНИЕ",
        "о взыскании задолженности",
        "",
        "Между сторонами заключён **договор подряда** № 5.",
        "Ответчик обязательство не исполнил:",
        "- акт не подписан;",
        "- оплата не произведена.",
        "",
        "ПРОШУ:",
        "1. Взыскать долг 100 000 руб.",
    ])
    out = tmp_path / "isk.docx"
    res = md_build_content(md, out, tmp_path)
    assert out.exists()
    assert res["blocks_rendered"] > 0


def md_build_content(md: str, out: Path, tmp: Path) -> dict:
    md_file = tmp / "in.md"
    md_file.write_text(md, encoding="utf-8")
    return md_build(str(md_file), str(out), "pytest")


def test_md_build_without_title_fails_closed(tmp_path):
    md = "Просто текст без титула и шапки.\nВторая строка."
    out = tmp_path / "no_title.docx"
    try:
        md_build_content(md, out, tmp_path)
    except MdBuildError as e:
        assert "fail-closed" in str(e)
    else:
        raise AssertionError("должна быть ошибка: титул не распознан")


def test_md_build_empty_fails_closed(tmp_path):
    import pytest
    with pytest.raises(MdBuildError, match="empty"):
        md_build_content("", tmp_path / "x.docx", tmp_path)


def test_md_bold_rendered(tmp_path):
    md = "В суд\nИстец: X\n\nИСКОВОЕ ЗАЯВЛЕНИЕ\n\nОплата **не произведена** по договору."
    out = tmp_path / "b.docx"
    md_build_content(md, out, tmp_path)
    doc = Document(str(out))
    bold_texts = [r.text for p in doc.paragraphs for r in p.runs if r.bold and r.text]
    assert any("не произведена" in t for t in bold_texts)


def test_normalize_fixes_fonts(tmp_path):
    # документ с «мусорным» шрифтом -> normalize -> TNR
    import docx as dx
    from docx.shared import Pt
    d = dx.Document()
    r = d.add_paragraph("Обычный абзац Calibri")
    r.runs[0].font.name = "Calibri"
    r.runs[0].font.size = Pt(12)
    f = tmp_path / "mixed.docx"
    d.save(str(f))
    from docx_cli.fonts_normalize import normalize_document
    res = normalize_document(str(f))
    assert res["changes"] >= 1
    d2 = Document(str(f))
    assert d2.paragraphs[0].runs[0].font.name == "Times New Roman"


def test_isk_template_updated_per_legal_review():
    """Рекомендации юриста 20.10: 333.21 НК в тексте, доверенность/подписант."""
    spec = json.loads((ROOT / "templates" / "iskovoe_zayavlenie.json").read_text(encoding="utf-8-sig"))
    body_text = json.dumps(spec, ensure_ascii=False)
    assert "333.21" in body_text            # правовое основание пошлины
    assert "доверенность" in body_text      # полномочия подписанта (ст. 61-62 АПК)
    assert "представителя" in body_text     # реквизиты подписанта


def test_templates_all_valid_specs():
    """Все шаблоны каталога валидны как DocSpec (extra=forbid строгость)."""
    from docx_cli.spec import DocSpec
    tdir = ROOT / "templates"
    files = list(tdir.glob("*.json"))
    assert len(files) >= 3
    for f in files:
        d = json.loads(f.read_text(encoding="utf-8-sig"))
        d = _fill(d)
        d["output"]["file_name"] = "t.docx"
        DocSpec.model_validate(d)


def _fill(o):
    if isinstance(o, dict):
        return {k: _fill(v) for k, v in o.items()}
    if isinstance(o, list):
        return [_fill(x) for x in o]
    if isinstance(o, str):
        return re.sub(r"\{[а-яё_]+\}", "DEMO", o)
    return o
