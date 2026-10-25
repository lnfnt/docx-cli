"""docx-cli v0.5.0 приёмка: profile-extract/validate/repair, agent-read, verify-forbidden.

Clean-room реализация профиль-цикла (extract/validate/repair), agent-read
и verify-forbidden. Негативки обязательны.
"""
import json
import subprocess
import sys
from pathlib import Path

import docx as dx
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx_cli.profiles import extract, load_profile, repair, validate  # noqa: E402
from docx_cli.agent_reader import read  # noqa: E402

import tempfile

_SMOKE_TMP = Path(tempfile.mkdtemp(prefix="docx_smoke_"))
_SMOKE_MD = '''В Арбитражный суд города Москвы
Истец: ООО Тест, ИНН 7701001001, ОГРН 1157700000001
Ответчик: ООО Ответчик, ИНН 7702002002, ОГРН 1167700000002
Дело № А40-999/2026

ИСКОВОЕ ЗАЯВЛЕНИЕ
о взыскании задолженности

Между сторонами заключен **договор подряда** № 5.
Ответчик обязательство не исполнил:
- акт не подписан;
- оплата не произведена.

ПРОШУ:
1. Взыскать долг 100 000 руб.'''


def _synth_smoke_isk() -> str:
    """Синтетический образец: md-build искового (court-профиль) вместо внешнего /tmp-артефакта."""
    from docx_cli.md_court import md_build

    md = _SMOKE_TMP / "isk.md"
    md.write_text(_SMOKE_MD, encoding="utf-8")
    out = _SMOKE_TMP / "smoke_isk.docx"
    md_build(str(md), str(out), "pytest")
    return str(out)


SMOKE = _synth_smoke_isk()


def test_extract_from_sample():
    prof = extract(SMOKE)
    assert prof['font']['name'].lower().startswith('times')
    assert prof['margins']['left'] > 2  # суд: левое ~3
    assert prof['margins']['right'] < 2


def test_validate_clean_and_dirty(tmp_path):
    prof = extract(SMOKE)
    pj = tmp_path / 'p.json'
    pj.write_text(json.dumps(prof, ensure_ascii=False), encoding='utf-8')
    assert validate(SMOKE, load_profile(str(pj))) == []  # образец чист против своего профиля
    # грязный: Calibri
    d = dx.Document()
    d.add_paragraph('Текст Calibri').runs[0].font.name = 'Calibri'
    bad = tmp_path / 'bad.docx'
    d.save(str(bad))
    problems = validate(str(bad), load_profile(str(pj)))
    assert any(p['param'] in ('font_name', 'run_font_name') for p in problems)


def test_repair_before_after(tmp_path):
    prof = extract(SMOKE)
    pj = tmp_path / 'p.json'
    pj.write_text(json.dumps(prof, ensure_ascii=False), encoding='utf-8')
    d = dx.Document()
    d.add_paragraph('Кривой документ Calibri')
    bad = tmp_path / 'bad.docx'
    d.save(str(bad))
    res = repair(str(bad), load_profile(str(pj)), tmp_path / 'fixed.docx')
    assert res['before'] > 0 and res['after'] == 0 and res['fixed'] == res['before']
    fixed = dx.Document(str(tmp_path / 'fixed.docx'))
    assert fixed.styles['Normal'].font.name.lower().startswith('times')


def test_agent_read_window_and_cursor(tmp_path):
    res = read(SMOKE, 0, 200)
    assert res['ok' if False else 'total_chars'] > 0
    if res['next_cursor'] is not None:
        res2 = read(SMOKE, res['next_cursor'], 200)
        # окна не пересекаются: первый блок второго окна начинается после конца первого
        assert res2['window'][0] == res['next_cursor']


def test_agent_read_missing_file(tmp_path):
    from docx_cli.agent_reader import AgentReadError
    try:
        read(str(tmp_path / 'nope.docx'))
    except AgentReadError:
        pass
    else:
        raise AssertionError('должна быть ошибка')


def test_verify_forbidden_hits_and_clean(tmp_path):
    d = dx.Document()
    d.add_paragraph('Согласно ч. 18 ст. 36 закона...')
    d.add_paragraph('Чистый абзац.')
    bad = tmp_path / 'doc.docx'
    d.save(str(bad))
    fl = tmp_path / 'forbid.txt'
    fl.write_text('# комментарии игнорируются\nч. 18 ст. 36\n', encoding='utf-8')
    r = subprocess.run([sys.executable, '-m', 'docx_cli.cli', 'verify-forbidden',
                        str(bad), '--forbid-file', str(fl)],
                       capture_output=True, text=True)
    out = json.loads(r.stdout)
    assert r.returncode == 1 and out['status'] == 'forbidden_found'
    assert out['hits'][0]['pattern'] == 'ч. 18 ст. 36'
    # чистый документ -> exit 0
    clean = tmp_path / 'clean.docx'
    d2 = dx.Document()
    d2.add_paragraph('Согласно ч. 9 ст. 36 закона...')
    d2.save(str(clean))
    r2 = subprocess.run([sys.executable, '-m', 'docx_cli.cli', 'verify-forbidden',
                         str(clean), '--forbid-file', str(fl)],
                        capture_output=True, text=True)
    assert r2.returncode == 0 and json.loads(r2.stdout)['status'] == 'ok'
