"""docx-cli v0.7.0 §3 B-05/B-06/B-08 vectors (audit F8b/F7/F6-md)."""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from docx_cli.md_court import _normalize_dashes, md_build
from docx_cli.validate import validate_document


def _md(tmp_path, body):
    md = tmp_path / "src.md"
    md.write_text(body, encoding="utf-8")
    out = tmp_path / "out.docx"
    return md, out


def test_v070_f8b_heading_number_not_leaked(tmp_path):
    md, out = _md(tmp_path, "ИСКОВОЕ ЗАЯВЛЕНИЕ\nИстец: Иванов\n\n## 1. Нарушение норм\nТекст.\n")
    md_build(str(md), str(out), "pytest")
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert not any("##" in t for t in texts)
    assert any("1. Нарушение норм" in t for t in texts)


def test_v070_f8b_plain_heading_centered(tmp_path):
    md, out = _md(tmp_path, "ИСКОВОЕ ЗАЯВЛЕНИЕ\nИстец: Иванов\n\n## Обстоятельства дела\nТекст.\n")
    md_build(str(md), str(out), "pytest")
    doc = Document(str(out))
    p = next(p for p in doc.paragraphs if "Обстоятельства дела" in p.text)
    assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(r.bold for r in p.runs if r.text)


def test_v070_f7_dash_normalized_everywhere(tmp_path):
    body = ("ИСКОВОЕ ЗАЯВЛЕНИЕ\nИстец: Иванов – И.И.\n\n"
            "## Заголовок – с дашем\n\nТело – с дашем.\n"
            "Приложения:\n1. Копия – 1 экз.\n")
    md, out = _md(tmp_path, body)
    res = md_build(str(md), str(out), "pytest")
    doc = Document(str(out))
    texts = "".join(p.text for p in doc.paragraphs)
    assert "\u2013" not in texts
    assert "—" in texts
    assert res["report"] and any("dash" in str(x).lower() or "dash_normalized" in str(x) for x in res["report"])


def test_v070_f7_dash_report_records_fix(tmp_path):
    md, out = _md(tmp_path, "ИСКОВОЕ ЗАЯВЛЕНИЕ\nИстец: Иванов\n\nТело – с дашем.\n")
    res = md_build(str(md), str(out), "pytest")
    assert res.get("dash_normalized", 0) >= 1


def test_v070_f6md_hanging_indent_for_ten(tmp_path):
    md, out = _md(tmp_path, "ИСКОВОЕ ЗАЯВЛЕНИЕ\nИстец: Иванов\n\n10. десятый пункт\n")
    md_build(str(md), str(out), "pytest")
    doc = Document(str(out))
    p = next(p for p in doc.paragraphs if "десятый" in p.text)
    assert abs(p.paragraph_format.left_indent.cm - 1.25) <= 0.05
    assert p.paragraph_format.first_line_indent.cm <= -0.9


def test_v070_f6md_reflow_twelve(tmp_path):
    lines = ["ИСКОВОЕ ЗАЯВЛЕНИЕ", "Истец: Иванов", ""]
    for i in range(1, 13):
        lines.append(f"{i}. пункт {i}")
    md, out = _md(tmp_path, "\n".join(lines) + "\n")
    md_build(str(md), str(out), "pytest")
    res = validate_document(str(out))
    assert res["errors"] == [], res["errors"]


def test_v070_f7_normalize_dashes_function():
    assert _normalize_dashes("a – b") == "a — b"
