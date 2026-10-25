"""docx-cli v0.7.0 §4 B-04 vectors (audit F8)."""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_cli.md_court import md_build


def _md(tmp_path, body):
    md = tmp_path / "src.md"
    md.write_text(body, encoding="utf-8")
    out = tmp_path / "out.docx"
    return md, out


def test_v070_f8_title_priority_requisite_block(tmp_path):
    """B1/A3: «Исковое заявление № 42» в блоке реквизитов не титул."""
    body = ("Истец: Иванов\nОтветчик: Петров\nИсковое заявление № 42\n\n"
            "ОТЗЫВ ИСТЦА\nна возражения\nТекст.\n")
    md, out = _md(tmp_path, body)
    md_build(str(md), str(out), "pytest")
    doc = Document(str(out))
    titles = [p for p in doc.paragraphs
              if p.alignment == WD_ALIGN_PARAGRAPH.CENTER
              and any(r.bold for r in p.runs if r.text)]
    assert len(titles) == 1
    assert "ОТЗЫВ ИСТЦА" in titles[0].text
    req = next(p for p in doc.paragraphs if "Исковое заявление № 42" in p.text)
    assert req.alignment != WD_ALIGN_PARAGRAPH.CENTER


def test_v070_f8_title_candidate_in_requisites_not_centered(tmp_path):
    """B2: строка-кандидат в реквизитах не получает титульное форматирование."""
    body = ("В Арбитражный суд города Москвы\n"
            "Истец: ООО Тест\n"
            "Заявление № 42\n\n"
            "ИСКОВОЕ ЗАЯВЛЕНИЕ\nТекст.\n")
    md, out = _md(tmp_path, body)
    md_build(str(md), str(out), "pytest")
    doc = Document(str(out))
    req = next(p for p in doc.paragraphs if "Заявление № 42" in p.text)
    assert not any(r.bold and r.font.size and r.font.size.pt == 14 for r in req.runs if r.text)


def test_v070_f8_single_clean_title_recognized(tmp_path):
    """B3: ровно один чистый кандидат становится титулом."""
    body = ("Истец: Иванов\n\nИСКОВОЕ ЗАЯВЛЕНИЕ\nТекст.\n")
    md, out = _md(tmp_path, body)
    md_build(str(md), str(out), "pytest")
    doc = Document(str(out))
    titles = [p for p in doc.paragraphs if "ИСКОВОЕ ЗАЯВЛЕНИЕ" in p.text]
    assert titles and titles[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
