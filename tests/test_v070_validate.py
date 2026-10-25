"""docx-cli v0.7.0 §1 V-01/V-02 vectors (audit F5).

Синтетические .docx через python-docx в tmp_path. Векторы B1-B4 + skip-канал.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from docx_cli.validate import (
    _check_body_formatting,
    _classify_paragraphs,
    _is_service_signature,
    validate_document,
)


def _svc_doc():
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Защитник Иванов И.И.")
    return doc


def test_v070_f5_b1_service_signature_no_first_line_indent_required():
    """B1: подпись пробелами без красной строки не требует first_line_indent."""
    doc = _svc_doc()
    errors = _check_body_formatting(doc, None)
    assert not any("first_line_indent" in e for e in errors)


def test_v070_f5_b2_service_signature_alignment_still_checked():
    """B2: подпись пробелами проверяется на alignment."""
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Защитник Иванов И.И.")
    errors = _check_body_formatting(doc, None)
    assert any("alignment" in e for e in errors)


def test_v070_f5_b3_service_signature_line_spacing_checked():
    """B3: подпись пробелами проверяется на интервал (line_spacing)."""
    from docx.enum.text import WD_LINE_SPACING
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.add_run("Защитник Иванов И.И.")
    errors = _check_body_formatting(doc, None)
    assert any("line_spacing" in e for e in errors)


def test_v070_f5_b4_appendices_fallback_stops_at_empty_paragraph():
    """B4: фолбэк «Приложения:» снимает проверки только до пустого параграфа."""
    doc = Document()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; t.add_run("ЗАГОЛОВОК")
    b = doc.add_paragraph(); b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    b.paragraph_format.first_line_indent = Cm(1.25); b.add_run("Текст.")
    a = doc.add_paragraph(); a.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    a.paragraph_format.first_line_indent = Cm(1.25); a.add_run("Приложения:")
    o = doc.add_paragraph(); o.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    o.paragraph_format.first_line_indent = Cm(1.25); o.add_run("Ордер адвоката №1.")
    doc.add_paragraph()
    after = doc.add_paragraph(); after.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    after.paragraph_format.first_line_indent = Cm(1.25); after.add_run("Защитник указал, что продолжение.")
    roles = _classify_paragraphs(doc, None)
    idx_after = next(i for i, p in enumerate(doc.paragraphs) if "Защитник указал" in p.text)
    assert roles[idx_after] == "body"
    idx_order = next(i for i, p in enumerate(doc.paragraphs) if "Ордер" in p.text)
    assert roles[idx_order] != "body"


def test_v070_f5_skip_channel_service_signature(tmp_path):
    """skip-канал: служебная подпись даёт check_skipped запись."""
    doc = _svc_doc()
    skipped = []
    _classify_paragraphs(doc, None, skipped=skipped)
    assert any(f.get("type") == "check_skipped" for f in skipped)
    assert any("signature" in str(f.get("reason", "")) for f in skipped)


def test_v070_f5_skip_channel_validate_document(tmp_path):
    """skip-канал: validate_document возвращает skipped-список."""
    doc = _svc_doc()
    path = tmp_path / "svc.docx"
    doc.save(str(path))
    res = validate_document(str(path))
    assert "skipped" in res
    assert isinstance(res["skipped"], list)
    assert any(f.get("type") == "check_skipped" for f in res["skipped"])


def test_v070_f5_is_service_signature_narrowing():
    """_is_service_signature: пробельная подпись без tab+___."""
    doc = _svc_doc()
    p = doc.paragraphs[0]
    assert _is_service_signature(p, p.text.strip()) is True
    # proper signature with tab+___ is not service-signature heuristic
    doc2 = Document()
    p2 = doc2.add_paragraph()
    p2.add_run("Представитель истца\t___________ /И.И./")
    p2.paragraph_format.first_line_indent = Cm(0)
    assert _is_service_signature(p2, p2.text.strip()) is False
