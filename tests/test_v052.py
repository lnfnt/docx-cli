"""docx-cli v0.5.2 приёмка: inline-разметка <u>/_ без утечки тегов в текст
+ каноническая геометрия списков в md-build.

Дефекты зафиксированы внешним аудитом (п.3):
1. `<u>_…_</u>` в md протекал в docx как буквальный текст.
2. `- …` bullet-параграфы строились с first=0/left=0.6cm — валидатор не
   распознавал их как структурный список и браковал по first_line_indent.
"""
from pathlib import Path

import docx

from docx_cli.md_court import md_build
from docx_cli.validate import validate_document

MD_BASE = """# ДОПОЛНЕНИЯ (№ 2)
Кратко: {body}
"""

OUT = "out.docx"


def build(tmp_path: Path, body: str) -> Path:
    md = tmp_path / "src.md"
    md.write_text(MD_BASE.format(body=body), encoding="utf-8")
    out = tmp_path / OUT
    md_build(str(md), str(out), "pytest")
    return out


def test_underline_italic_no_tag_leak(tmp_path):
    out = build(tmp_path, "<u>_подчёркнутый курсив_</u> обычный")
    doc = docx.Document(str(out))
    body_paras = [p for p in doc.paragraphs if "обычный" in p.text]
    assert len(body_paras) == 1
    p = body_paras[0]
    assert "<u>" not in p.text and "</u>" not in p.text and "_" not in p.text
    marked = [r for r in p.runs if r.text == "подчёркнутый курсив"]
    assert marked and marked[0].underline is True and marked[0].italic is True
    plain = [r for r in p.runs if "обычный" in r.text]
    assert plain[0].underline is None and plain[0].italic is None


def test_bold_still_works(tmp_path):
    out = build(tmp_path, "**важный** абзац")
    doc = docx.Document(str(out))
    runs = [r for p in doc.paragraphs for r in p.runs if r.text == "важный"]
    assert runs and runs[0].bold is True
    assert "*" not in "".join(p.text for p in doc.paragraphs)


def test_combination_bold_underline(tmp_path):
    out = build(tmp_path, "**полужирный<u> и подчёркнутый</u>** конец")
    doc = docx.Document(str(out))
    runs = [r for p in doc.paragraphs for r in p.runs if "и подчёркнутый" in r.text]
    assert runs and runs[0].bold is True and runs[0].underline is True


def test_intraword_underscore_is_literal(tmp_path):
    out = build(tmp_path, "файл doc_creator.py на столе")
    doc = docx.Document(str(out))
    texts = "".join(p.text for p in doc.paragraphs)
    assert "doc_creator.py" in texts


def test_bullet_hanging_geometry_and_validate(tmp_path):
    md = tmp_path / "src.md"
    md.write_text(
        "# ДОПОЛНЕНИЯ (№ 2)\n"
        "Пункт вводный.\n"
        "- <u>_в нарушение ст. 9 закона;_</u>\n"
        "- второй пункт списка.\n",
        encoding="utf-8",
    )
    out = tmp_path / OUT
    md_build(str(md), str(out), "pytest")
    doc = docx.Document(str(out))
    bullet = [p for p in doc.paragraphs if p.text.startswith("в нарушение")][0]
    from docx.shared import Cm
    # python-docx хранит инденты в twips, обратное чтение даёт EMU с точностью
    # до ~220 (1.25cm: 450215 вместо 450000) — то же самое делает blocks.py.
    assert abs(bullet.paragraph_format.left_indent - Cm(1.25)) < 500
    assert abs(bullet.paragraph_format.first_line_indent + Cm(0.75)) < 500
    # сквозная проверка: собственный validate не бракует сборку
    res = validate_document(str(out))
    indent_errors = [e for e in res["errors"] if "first_line_indent" in e]
    assert indent_errors == [], indent_errors


def test_en_dash_remap_survives_inline_split(tmp_path):
    out = build(tmp_path, "<u>_отчёт 108–135 страниц_</u>")
    doc = docx.Document(str(out))
    texts = "".join(p.text for p in doc.paragraphs)
    assert "\u2013" not in texts and "108—135" in texts
