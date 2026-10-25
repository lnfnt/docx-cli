"""docx-cli v0.6.0 scan vectors (v0.6 §1-§4, audit F1/F2/F3).

Синтетические .docx через python-docx + ручная правка XML в zip (tmp_path,
живые файлы не трогаются). Каждый тест в docstring ссылается на F-номер
аудита 21.10 (F1/F2/F3).
"""
from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from docx_cli.inject_scan import scan_all


def _base_doc(tmp_path: Path, name: str = "base.docx") -> Path:
    doc = Document()
    doc.add_paragraph("X")
    path = tmp_path / name
    doc.save(str(path))
    return path


def _mutate(base: Path, out: Path, *, doc_transform=None, styles_transform=None,
            extra=None, drop_document=False, corrupt_document=False) -> None:
    zin = zipfile.ZipFile(base)
    names = zin.namelist()
    with zipfile.ZipFile(out, "w") as zout:
        for n in names:
            if drop_document and n == "word/document.xml":
                continue
            data = zin.read(n)
            if n == "word/document.xml":
                if corrupt_document:
                    data = b"<w:document><w:body><w:p><w:r><w:t>bad"
                elif doc_transform is not None:
                    data = doc_transform(data.decode("utf-8")).encode("utf-8")
            elif n == "word/styles.xml" and styles_transform is not None:
                data = styles_transform(data.decode("utf-8")).encode("utf-8")
            zout.writestr(n, data)
        for n, data in (extra or {}).items():
            zout.writestr(n, data)
    zin.close()


def test_v060_f1_vanish_via_rstyle(tmp_path: Path) -> None:
    """F1: vanish унаследованный через w:rStyle -> стиль с vanish."""
    base = _base_doc(tmp_path)
    out = tmp_path / "vanish_rstyle.docx"

    def styles(s: str) -> str:
        return s.replace("</w:styles>",
                        '<w:style w:type="character" w:styleId="Hidden">'
                        '<w:rPr><w:vanish/></w:rPr></w:style></w:styles>')

    def doc(s: str) -> str:
        return s.replace("<w:r><w:t>X</w:t></w:r>",
                        '<w:r><w:rPr><w:rStyle w:val="Hidden"/></w:rPr>'
                        '<w:t>secret vanish</w:t></w:r>')

    _mutate(base, out, doc_transform=doc, styles_transform=styles)
    res = scan_all(str(out))
    assert res["clean"] is False
    types = {f["type"] for f in res["findings"]}
    assert "hidden_text_vanish" in types
    f = next(x for x in res["findings"] if x["type"] == "hidden_text_vanish")
    assert "secret vanish" in f["text"]


def test_v060_f1_white_via_style_basedon(tmp_path: Path) -> None:
    """F1: белый текст через стиль + basedOn-цепочку."""
    base = _base_doc(tmp_path)
    out = tmp_path / "white_chain.docx"

    def styles(s: str) -> str:
        return s.replace("</w:styles>",
                        '<w:style w:type="character" w:styleId="Base">'
                        '<w:rPr><w:color w:val="ffffff"/><w:sz w:val="4"/></w:rPr></w:style>'
                        '<w:style w:type="character" w:styleId="Mid">'
                        '<w:basedOn w:val="Base"/><w:rPr><w:szCs w:val="4"/></w:rPr></w:style>'
                        '<w:style w:type="character" w:styleId="Top">'
                        '<w:basedOn w:val="Mid"/></w:style></w:styles>')

    def doc(s: str) -> str:
        return s.replace("<w:r><w:t>X</w:t></w:r>",
                        '<w:r><w:rPr><w:rStyle w:val="Top"/></w:rPr>'
                        '<w:t>secret white micro</w:t></w:r>')

    _mutate(base, out, doc_transform=doc, styles_transform=styles)
    res = scan_all(str(out))
    assert res["clean"] is False
    types = {f["type"] for f in res["findings"]}
    assert "white_text_on_light_bg" in types


def test_v060_f1_near_white_fffffe(tmp_path: Path) -> None:
    """F1/F2: near-white fffffe."""
    base = _base_doc(tmp_path)
    out = tmp_path / "near_fffffe.docx"

    def doc(s: str) -> str:
        return s.replace("<w:r><w:t>X</w:t></w:r>",
                        '<w:r><w:rPr><w:color w:val="fffffe"/></w:rPr>'
                        '<w:t>secret near</w:t></w:r>')

    _mutate(base, out, doc_transform=doc)
    res = scan_all(str(out))
    assert res["clean"] is False
    types = {f["type"] for f in res["findings"]}
    assert "near_white" in types


def test_v060_f1_near_white_f8f8f8(tmp_path: Path) -> None:
    """F1/F2: near-white F8F8F8."""
    base = _base_doc(tmp_path)
    out = tmp_path / "near_f8f8f8.docx"

    def doc(s: str) -> str:
        return s.replace("<w:r><w:t>X</w:t></w:r>",
                        '<w:r><w:rPr><w:color w:val="F8F8F8"/></w:rPr>'
                        '<w:t>secret near</w:t></w:r>')

    _mutate(base, out, doc_transform=doc)
    res = scan_all(str(out))
    assert res["clean"] is False
    types = {f["type"] for f in res["findings"]}
    assert "near_white" in types


def test_v060_f1_theme_color_light1(tmp_path: Path) -> None:
    """F1/F2: themeColor light1."""
    base = _base_doc(tmp_path)
    out = tmp_path / "theme_light1.docx"

    def doc(s: str) -> str:
        return s.replace("<w:r><w:t>X</w:t></w:r>",
                        '<w:r><w:rPr><w:color w:themeColor="light1"/></w:rPr>'
                        '<w:t>secret theme</w:t></w:r>')

    _mutate(base, out, doc_transform=doc)
    res = scan_all(str(out))
    assert res["clean"] is False
    types = {f["type"] for f in res["findings"]}
    assert "theme_white" in types


def test_v060_f1_micro_sz4(tmp_path: Path) -> None:
    """F1: микро sz=4 (<=2pt)."""
    base = _base_doc(tmp_path)
    out = tmp_path / "micro_sz4.docx"

    def doc(s: str) -> str:
        return s.replace("<w:r><w:t>X</w:t></w:r>",
                        '<w:r><w:rPr><w:sz w:val="4"/></w:rPr>'
                        '<w:t>secret micro</w:t></w:r>')

    _mutate(base, out, doc_transform=doc)
    res = scan_all(str(out))
    assert res["clean"] is False
    types = {f["type"] for f in res["findings"]}
    assert "microfont" in types
    f = next(x for x in res["findings"] if x["type"] == "microfont")
    assert f["size_pt"] == 2.0


def test_v060_f1_micro_szcs4(tmp_path: Path) -> None:
    """F1: микро szCs=4 (<=2pt)."""
    base = _base_doc(tmp_path)
    out = tmp_path / "micro_szcs4.docx"

    def doc(s: str) -> str:
        return s.replace("<w:r><w:t>X</w:t></w:r>",
                        '<w:r><w:rPr><w:szCs w:val="4"/></w:rPr>'
                        '<w:t>secret micro</w:t></w:r>')

    _mutate(base, out, doc_transform=doc)
    res = scan_all(str(out))
    assert res["clean"] is False
    types = {f["type"] for f in res["findings"]}
    assert "microfont" in types


def test_v060_f2_vanish_in_header(tmp_path: Path) -> None:
    """F2: vanish в header2.xml; findings path=часть."""
    base = _base_doc(tmp_path)
    out = tmp_path / "header_vanish.docx"
    header_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:r><w:rPr><w:vanish/></w:rPr><w:t>hidden header</w:t></w:r></w:p>'
        '</w:hdr>'
    )
    _mutate(base, out, extra={"word/header2.xml": header_xml})
    res = scan_all(str(out))
    assert res["clean"] is False
    f = next(x for x in res["findings"] if x["type"] == "hidden_text_vanish")
    assert "hidden header" in f["text"]
    assert f["path"] == "word/header2.xml"


def test_v060_f2_vanish_in_footer(tmp_path: Path) -> None:
    """F2: vanish в footer1.xml; findings path=часть."""
    base = _base_doc(tmp_path)
    out = tmp_path / "footer_vanish.docx"
    footer_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:r><w:rPr><w:vanish/></w:rPr><w:t>hidden footer</w:t></w:r></w:p>'
        '</w:ftr>'
    )
    _mutate(base, out, extra={"word/footer1.xml": footer_xml})
    res = scan_all(str(out))
    assert res["clean"] is False
    f = next(x for x in res["findings"] if x["type"] == "hidden_text_vanish")
    assert "hidden footer" in f["text"]
    assert f["path"] == "word/footer1.xml"


def test_v060_f2_vanish_in_footnote(tmp_path: Path) -> None:
    """F2: vanish в footnotes.xml; findings path=часть."""
    base = _base_doc(tmp_path)
    out = tmp_path / "footnote_vanish.docx"
    footnote_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:id="1"><w:p><w:r><w:rPr><w:vanish/></w:rPr>'
        '<w:t>hidden footnote</w:t></w:r></w:p></w:footnote>'
        '</w:footnotes>'
    )
    _mutate(base, out, extra={"word/footnotes.xml": footnote_xml})
    res = scan_all(str(out))
    assert res["clean"] is False
    f = next(x for x in res["findings"] if x["type"] == "hidden_text_vanish")
    assert "hidden footnote" in f["text"]
    assert f["path"] == "word/footnotes.xml"


def test_v060_f3_not_a_docx(tmp_path: Path) -> None:
    """F3: zip без word/document.xml -> not_a_docx finding, clean=False."""
    out = tmp_path / "no_document.docx"
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/styles.xml",
                   '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    res = scan_all(str(out))
    assert res["clean"] is False
    assert res["findings"][0]["type"] == "not_a_docx"
    assert res["findings"][0]["path"] == "word/document.xml"


def test_v060_f3_broken_xml(tmp_path: Path) -> None:
    """F3: битый XML -> xml_parse_error конверт без traceback."""
    base = _base_doc(tmp_path)
    out = tmp_path / "broken.docx"
    _mutate(base, out, corrupt_document=True)
    res = scan_all(str(out))
    assert res["clean"] is False
    assert res["findings"][0]["type"] == "xml_parse_error"


def test_v060_f3_missing_file(tmp_path: Path) -> None:
    """F3: отсутствующий файл -> io_error конверт."""
    missing = tmp_path / "nope.docx"
    res = scan_all(str(missing))
    assert res["clean"] is False
    assert res["findings"][0]["type"] == "io_error"


def test_v060_f1_clean_document(tmp_path: Path) -> None:
    """F1/F2: чистый документ -> clean=True."""
    doc = Document()
    doc.add_paragraph("Обычный текст без скрытых элементов.")
    path = tmp_path / "clean.docx"
    doc.save(str(path))
    res = scan_all(str(path))
    assert res["clean"] is True
    assert res["findings"] == []


def test_v060_f1_dark_text_not_flagged(tmp_path: Path) -> None:
    """F1/F2: обычный тёмный текст не флагается."""
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Обычный тёмный текст")
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    r.font.size = Pt(12)
    path = tmp_path / "dark.docx"
    doc.save(str(path))
    res = scan_all(str(path))
    assert res["clean"] is True
    assert res["findings"] == []
