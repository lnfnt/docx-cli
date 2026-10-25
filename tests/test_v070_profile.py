"""docx-cli v0.7.0 §5 P-01 vectors (audit F9).

Синтетические .docx через python-docx в tmp_path. Профиль-цикл честный:
font только факт, обход таблиц/колонтитулов, validate/repair покрывают
header/footer/ячейки.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt

from docx_cli.profiles import extract, load_profile, repair, validate


def _sample_tnr(tmp_path: Path, name: str = "sample.docx") -> Path:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    p = doc.add_paragraph("Обычный текст образца.")
    p.runs[0].font.name = "Times New Roman"
    path = tmp_path / name
    doc.save(str(path))
    return path


def test_v070_f9_extract_font_fact_no_fallback(tmp_path: Path):
    """F9: нет явного шрифта -> name None + guessed=True (без TNR-фолбэка)."""
    doc = Document()
    doc.add_paragraph("Текст без явного шрифта.")
    path = tmp_path / "nofont.docx"
    doc.save(str(path))
    prof = extract(str(path))
    assert prof["font"]["name"] is None
    assert prof["font"].get("guessed") is True


def test_v070_f9_extract_font_fact_from_run(tmp_path: Path):
    """F9: шрифт берётся фактом из run rPr."""
    doc = Document()
    p = doc.add_paragraph("Текст")
    p.runs[0].font.name = "Calibri"
    path = tmp_path / "runfont.docx"
    doc.save(str(path))
    prof = extract(str(path))
    assert prof["font"]["name"] == "Calibri"
    assert prof["font"].get("guessed") is False


def test_v070_f9_extract_includes_header_footer_table(tmp_path: Path):
    """F9: extract обходит колонтитулы и таблицы."""
    doc = Document()
    doc.sections[0].footer.paragraphs[0].add_run("подвал").font.name = "Arial"
    path = tmp_path / "parts.docx"
    doc.save(str(path))
    prof = extract(str(path))
    assert prof["font"]["name"] == "Arial"
    assert prof["font"].get("guessed") is False


def test_v070_f9_validate_covers_header_footer_cells(tmp_path: Path):
    """F9: validate ловит run_font в header/footer/ячейках."""
    sample = _sample_tnr(tmp_path)
    prof = extract(str(sample))
    pj = tmp_path / "p.json"
    pj.write_text(json.dumps(prof, ensure_ascii=False), encoding="utf-8")
    profile = load_profile(str(pj))

    doc = Document()
    doc.sections[0].header.paragraphs[0].add_run("заголовок").font.name = "Calibri"
    doc.sections[0].footer.paragraphs[0].add_run("подвал").font.name = "Calibri"
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].paragraphs[0].add_run("ячейка").font.name = "Calibri"
    bad = tmp_path / "bad.docx"
    doc.save(str(bad))

    problems = validate(str(bad), profile)
    wheres = " ".join(p["where"] for p in problems)
    assert "header" in wheres.lower() or "Header" in wheres
    assert "footer" in wheres.lower() or "Footer" in wheres
    assert "table" in wheres.lower() or "cell" in wheres.lower()


def test_v070_f9_repair_covers_header_footer_cells(tmp_path: Path):
    """F9: repair чинит run_font в header/footer/ячейках."""
    sample = _sample_tnr(tmp_path)
    prof = extract(str(sample))
    pj = tmp_path / "p.json"
    pj.write_text(json.dumps(prof, ensure_ascii=False), encoding="utf-8")
    profile = load_profile(str(pj))

    doc = Document()
    doc.sections[0].header.paragraphs[0].add_run("заголовок").font.name = "Calibri"
    doc.sections[0].footer.paragraphs[0].add_run("подвал").font.name = "Calibri"
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].paragraphs[0].add_run("ячейка").font.name = "Calibri"
    bad = tmp_path / "bad.docx"
    doc.save(str(bad))

    out = tmp_path / "fixed.docx"
    res = repair(str(bad), profile, str(out))
    assert res["after"] == 0

    fixed = Document(str(out))
    for part in (fixed.sections[0].header, fixed.sections[0].footer):
        for p in part.paragraphs:
            for r in p.runs:
                assert (r.font.name or "").lower().startswith("times")
    for t in fixed.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        assert (r.font.name or "").lower().startswith("times")
