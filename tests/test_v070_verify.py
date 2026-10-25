"""docx-cli v0.7.0 §6 P-02 vectors (audit F10).

verify-forbidden читает те же части, что scan (headers/footers/footnotes/
endnotes); отсутствующий/битый файл = конверт io_error/xml_parse_error.
"""
from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document


def test_v070_f10_verify_forbidden_footer(tmp_path: Path):
    """F10: запрещённая строка в footer ловится."""
    doc = Document()
    doc.sections[0].footer.paragraphs[0].add_run("SECRET-FORBIDDEN")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    r = subprocess.run(
        [sys.executable, "-m", "docx_cli.cli", "verify-forbidden",
         str(path), "--forbid", "SECRET-FORBIDDEN"],
        capture_output=True, text=True,
    )
    out = json.loads(r.stdout)
    assert r.returncode == 1
    assert out["status"] in ("forbidden_found", "ok") is False or out["count"] >= 1
    assert out["count"] >= 1


def test_v070_f10_verify_forbidden_missing_file(tmp_path: Path):
    missing = tmp_path / "nope.docx"
    r = subprocess.run(
        [sys.executable, "-m", "docx_cli.cli", "verify-forbidden",
         str(missing), "--forbid", "x"],
        capture_output=True, text=True,
    )
    assert r.returncode == 1
    out = json.loads(r.stdout)
    assert "io_error" in json.dumps(out)


def test_v070_f10_verify_forbidden_broken_xml(tmp_path: Path):
    base = tmp_path / "base.docx"
    _d = Document(); _d.add_paragraph("X"); _d.save(str(base))
    # corrupt document.xml
    out = tmp_path / "broken.docx"
    import zipfile as zf
    zin = zf.ZipFile(base)
    with zf.ZipFile(out, "w") as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "word/document.xml":
                data = b"<w:document><w:body><w:p><w:r><w:t>bad"
            zout.writestr(n, data)
    zin.close()
    r = subprocess.run(
        [sys.executable, "-m", "docx_cli.cli", "verify-forbidden",
         str(out), "--forbid", "x"],
        capture_output=True, text=True,
    )
    assert r.returncode == 1
    out = json.loads(r.stdout)
    assert "xml_parse_error" in json.dumps(out)
