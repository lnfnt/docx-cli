"""docx-cli v0.5.1 регрессия по внешнему рецензированию.

Покрываю три подтверждённых дефекта v0.5.0:
  1) TITLE_WORDS: опечатка «АПЕЛЯЦИОННАЯ», отсутствие вариаций
     «ДОПОЛНЕНИЯ (№ 2)» / «ДОПОЛНИТЕЛЬНАЯ АПЕЛЛЯЦИОННАЯ ЖАЛОБА» -> md-build
     падал fail-closed на рабочем md;
  2) validate: ложноположительный first_line_indent на блоке приложений и
     подписи, набранной пробелами (п. 59 «Ордер адвоката…», п. 61 «Защитник…»);
  3) CLI: agent-read не принимал --max-chars; verify-forbidden --forbid не
     принимал несколько значений за один флаг.
"""
from __future__ import annotations

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from docx_cli.cli import build_parser
from docx_cli.md_court import _is_title, md_build
from docx_cli.validate import _classify_paragraphs


# --- 1. титулы -------------------------------------------------------------

def test_title_correct_spelling_recognized():
    assert _is_title("АПЕЛЛЯЦИОННАЯ ЖАЛОБА")
    assert _is_title("Апелляционная жалоба")  # регистронезависимо


def test_title_legacy_typo_alias_kept():
    # документы с опечаткой в титуле не должны падать
    assert _is_title("АПЕЛЯЦИОННАЯ ЖАЛОБА")


def test_title_dopолнeniya_variants():
    assert _is_title("ДОПОЛНЕНИЯ (№ 2)")
    assert _is_title("ДОПОЛНЕНИЕ")
    assert _is_title("ДОПОЛНЕНИЯ")
    assert _is_title("ДОПОЛНЕНИЯ К АПЕЛЛЯЦИОННОЙ ЖАЛОБЕ")
    assert _is_title("ДОПОЛНЕНИЕ К ЖАЛОБЕ")
    assert _is_title("ДОПОЛНИТЕЛЬНАЯ АПЕЛЛЯЦИОННАЯ ЖАЛОБА")
    assert _is_title("ДОПОЛНЕНИЯ К АПЕЛЛЯЦИОННОЙ ЖАЛОБЕ (№ 2)")


def test_title_unknown_fail_closed_preserved():
    assert not _is_title("СОБРАНИЕ АКЦИОНЕРОВ")
    assert not _is_title("ПРОТОКОЛ ОСМОТРА")


def test_md_build_repro_h1_addendum_title(tmp_path):
    """Кейс из приёмки: первая строка md — «# ДОПОЛНЕНИЯ (№ 2)» (H1 + номер) —
    в v0.5.0 роняла md-build с «титульная строка не распознана»."""
    md = tmp_path / "dop.md"
    md.write_text(
        "# ДОПОЛНЕНИЯ (№ 2)\n"
        "к апелляционной жалобе на приговор.\n"
        "Защитник: адвокат Т.\n"
        "В производстве районного суда.\n"
        "ПРОШУ:\n"
        "1. Приговор отменить.\n",
        encoding="utf-8",
    )
    out = tmp_path / "out.docx"
    res = md_build(str(md), str(out), "pytest")
    assert out.is_file()
    assert res["blocks_rendered"] > 0
    import docx as _docx
    built = _docx.Document(str(out))
    title_p = next(p for p in built.paragraphs if p.text.strip())
    assert "#" not in title_p.text  # md-маркер не должен попасть в титул


# --- 2. служебный хвост не body --------------------------------------------

def _mk_doc(with_apps_block: bool):
    d = docx.Document()

    def add(text, align, first=None):
        p = d.add_paragraph()
        p.alignment = align
        if first is not None:
            p.paragraph_format.first_line_indent = first
        p.add_run(text)
        return p

    add("АПЕЛЛЯЦИОННАЯ ЖАЛОБА", WD_ALIGN_PARAGRAPH.CENTER)
    add("Обстоятельства дела следующие.", WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(1.25))
    if with_apps_block:
        add("Приложения:", WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(1.25))
        add("1. Копия для прокурора — на 2 экз.;", WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(1.25))
        add("Ордер адвоката № 2/1551 от 05.06.2026.", WD_ALIGN_PARAGRAPH.JUSTIFY, None)
        add("Защитник                                Иванов И.И.",
            WD_ALIGN_PARAGRAPH.JUSTIFY, None)
    else:
        add("Защитник                                Иванов И.И.",
            WD_ALIGN_PARAGRAPH.JUSTIFY, None)
    return d


def test_attachments_tail_not_body():
    roles = _classify_paragraphs(_mk_doc(True), None)
    assert roles.get(1) == "body"        # обычный body остался body
    assert roles.get(2) == "other"       # «Приложения:»
    assert roles.get(4) != "body"        # «Ордер адвоката…» (п.59 кейса приёмки)
    assert roles.get(5) != "body"        # подпись пробелами (п.61 кейса приёмки)


def test_signature_before_apps_block():
    roles = _classify_paragraphs(_mk_doc(False), None)
    assert roles.get(2) == "signature"   # «Защитник …» без красной строки
    assert roles[1] == "body"


def test_body_sentence_starting_zashchitnik_stays_body():
    """«Защитник указал, что…» с красной строкой — это body, не подпись."""
    d = docx.Document()
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("ХОДАТАЙСТВО")
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.add_run("Защитник указал, что доказательства получены с нарушением.")
    roles = _classify_paragraphs(d, None)
    assert roles.get(1) == "body"


# --- 3. CLI-согласованность --------------------------------------------------

def test_md_build_validate_roundtrip(tmp_path):
    """v0.5.1: сборка md-build обязана проходить собственный validate
    (в v0.5.0 била right_margin=1.0 и run-и без имени шрифта)."""
    md = tmp_path / "rt.md"
    md.write_text(
        "# АПЕЛЛЯЦИОННАЯ ЖАЛОБА\n"
        "на приговор районного суда.\n"
        "Считаю приговор незаконным.\n"
        "ПРОШУ:\n"
        "1. Приговор отменить.\n"
        "Приложения:\n"
        "1. Копия жалобы — на 2 экз.\n"
        "2. Ордер – на 1 экз.\n",
        encoding="utf-8",
    )
    out = tmp_path / "rt.docx"
    md_build(str(md), str(out), "pytest")
    from docx_cli.validate import validate_document
    res = validate_document(str(out))
    assert not res["errors"], res["errors"]
    assert res["profile"] == "gost", res["profile"]  # доминирующий 12pt угадан
    # en dash в источнике нормализован в em dash (no_en_dash)
    built = docx.Document(str(out))
    assert all("\u2013" not in p.text for p in built.paragraphs)


def test_agent_read_accepts_max_chars_alias():
    parser = build_parser()
    a = parser.parse_args(["agent-read", "f.docx", "--max-chars", "1000"])
    assert a.max_chars == 1000
    b = parser.parse_args(["agent-read", "f.docx", "--max", "700"])
    assert b.max_chars == 700


def test_verify_forbidden_multiple_values_per_flag():
    parser = build_parser()
    c = parser.parse_args(
        ["verify-forbidden", "f.docx", "--forbid", "А", "Б", "--forbid", "В"])
    assert c.forbid == [["А", "Б"], ["В"]]
