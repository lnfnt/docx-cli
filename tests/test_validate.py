"""Tests for built-in validation checks."""
from __future__ import annotations

import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_cli.builder import build_document
from docx_cli.spec import DocSpec
from docx_cli.validate import validate_document


class TestValidationChecks:
    def test_valid_document_passes(self, built_path, minimal_spec):
        result = validate_document(str(built_path), minimal_spec)
        assert not result["errors"]
        assert result["checks"]["margins"] == "pass"
        assert result["checks"]["fonts"] == "pass"
        assert result["checks"]["body_formatting"] == "pass"

    def test_placeholder_warning(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "[УКАЗАТЬ сумму]", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "ph.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "degraded"
        assert any("placeholder_present" in w for w in result["metadata"]["warnings"])

    def test_yo_warning(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "Текст с буквой ё", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "yo.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "degraded"
        assert any("yo_present" in w for w in result["metadata"]["warnings"])

    def test_yo_whitelist_requisite_header_not_warned(self, tmp_workspace, minimal_spec_dict):
        # Requisite header line (contains INN digits) may contain yo
        minimal_spec_dict["header"]["lines"] = ["ИНН 7701234567, ё"]
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "Текст", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "yo_whitelist.docx"
        result = build_document(spec, str(out))
        assert not any("yo_present" in w for w in result["metadata"]["warnings"])

    def test_yo_in_body_text_warns(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "счёт", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "yo_body.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "degraded"
        assert any("yo_present" in w for w in result["metadata"]["warnings"])

    def test_table_cell_en_dash_error(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {
                "type": "table",
                "headers": ["А"],
                "rows": [["Текст – с тире"]],
            }
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table_endash.docx"
        result = build_document(spec, str(out))
        assert result["ok"] is False
        assert result["status"] == "error"
        assert not out.exists()
        assert any("en_dash" in e for e in result["metadata"]["errors"])

    def test_table_cell_yo_warning(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {
                "type": "table",
                "headers": ["А"],
                "rows": [["Текст с ё"]],
            }
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table_yo.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "degraded"
        assert any("yo_present" in w for w in result["metadata"]["warnings"])

    def test_table_cell_placeholder_warning(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {
                "type": "table",
                "headers": ["А"],
                "rows": [["[УКАЗАТЬ сумму]"]],
            }
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "table_ph.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "degraded"
        assert any("placeholder_present" in w for w in result["metadata"]["warnings"])

    def test_font_name_none_error(self, tmp_workspace):
        out = tmp_workspace / "none_font.docx"
        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run("text without explicit font")
        # Leave run.font.name as None (inherited from style)
        assert run.font.name is None
        doc.save(str(out))
        result = validate_document(str(out))
        assert any("font_name" in e and "None" in e for e in result["errors"])

    def test_body_alignment_left_error(self, tmp_workspace):
        out = tmp_workspace / "align_left.docx"
        doc = Document()
        title = doc.add_paragraph("ЗАГОЛОВОК")
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        body = doc.add_paragraph("Обычный абзац")
        body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        doc.save(str(out))
        result = validate_document(str(out))
        assert any("alignment" in e for e in result["errors"])

    def test_numbered_list_out_of_order_single_error(self, tmp_workspace):
        out = tmp_workspace / "order.docx"
        doc = Document()
        for text in ("1. first", "3. third", "2. second"):
            doc.add_paragraph(text)
        doc.save(str(out))
        result = validate_document(str(out))
        seq_errors = [e for e in result["errors"] if "numbered_list_sequence" in e]
        assert len(seq_errors) == 1
        assert any("out of order" in e for e in result["errors"])

    def test_en_dash_error(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "Текст – с тире", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "endash.docx"
        result = build_document(spec, str(out))
        assert result["ok"] is False
        assert result["status"] == "error"
        assert not out.exists()
        assert any("en_dash" in e for e in result["metadata"]["errors"])

    def test_en_dash_warning_when_allowed(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "Текст – с тире", "inline_markdown": False}
        ]
        minimal_spec_dict["checks"]["no_en_dash"] = False
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "endash_warn.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "degraded"
        assert out.exists()
        assert any("en_dash" in w for w in result["metadata"]["warnings"])

    def test_numbered_list_gap_error(self, tmp_workspace):
        from docx import Document

        out = tmp_workspace / "gap.docx"
        doc = Document()
        doc.add_paragraph("1. first")
        doc.add_paragraph("3. third")
        doc.save(str(out))
        result = validate_document(str(out))
        assert any("numbered_list" in e for e in result["errors"])

    def test_numbered_list_duplicate_is_restart_ok_v053(self, tmp_workspace):
        """v0.5.3 (фикс 3 — согласованная смена семантики):
        независимые перечни [1 | 1] легитимны — рестарт группы при n==1, если
        предыдущая группа полная любой длины, включая одноэлементную [1].
        В v0.5.2 это ложно браковалось как «out of order»."""
        from docx import Document

        out = tmp_workspace / "dup.docx"
        doc = Document()
        doc.add_paragraph("1. first")
        doc.add_paragraph("1. duplicate")
        doc.save(str(out))
        result = validate_document(str(out))
        assert not any("numbered_list" in e for e in result["errors"])

    def test_numbered_list_valid_sequence(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "numbered_list", "items": ["a", "b", "c"]}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "seq.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "ok"

    def test_norm_refs_warning(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["checks"]["require_norms_refs"] = True
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "norms.docx"
        result = build_document(spec, str(out))
        assert result["status"] == "degraded"
        assert any("norm_refs_missing" in w for w in result["metadata"]["warnings"])

    def test_norm_refs_present_no_warning(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["checks"]["require_norms_refs"] = True
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "Согласно ст. 309 ГК РФ", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "norms_ok.docx"
        result = build_document(spec, str(out))
        assert not any("norm_refs_missing" in w for w in result["metadata"]["warnings"])


class TestStandaloneValidation:
    def test_validate_ok(self, built_path):
        result = validate_document(str(built_path))
        assert result["checks"]["margins"] == "pass"
        assert result["checks"]["fonts"] == "pass"

    def test_validate_bad_file(self, tmp_workspace):
        bad = tmp_workspace / "not_a_docx.docx"
        bad.write_text("hello")
        result = validate_document(str(bad))
        assert result["errors"]

    def test_validate_missing_file(self, tmp_workspace):
        missing = tmp_workspace / "missing.docx"
        result = validate_document(str(missing))
        assert result["errors"]

    def test_standalone_validate_table_contents(self, tmp_workspace):
        out = tmp_workspace / "standalone_table.docx"
        doc = Document()
        title = doc.add_paragraph("ЗАГОЛОВОК")
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "Текст – ё [УКАЗАТЬ]"
        doc.save(str(out))
        result = validate_document(str(out))
        assert any("en_dash" in e for e in result["errors"])
        assert any("yo_present" in w for w in result["warnings"])
        assert any("placeholder_present" in w for w in result["warnings"])


class TestDeletionOnError:
    def test_partial_output_deleted_on_error(self, tmp_workspace, minimal_spec_dict):
        minimal_spec_dict["body"] = [
            {"type": "paragraph", "text": "ошибка – en dash", "inline_markdown": False}
        ]
        spec = DocSpec.from_dict(minimal_spec_dict)
        out = tmp_workspace / "should_be_deleted.docx"
        result = build_document(spec, str(out))
        assert result["ok"] is False
        assert not out.exists()

    def test_partial_output_deleted_on_unexpected_exception(
        self, tmp_workspace, minimal_spec, monkeypatch
    ):
        out = tmp_workspace / "exception_deleted.docx"

        def boom(*args, **kwargs):
            raise RuntimeError("validation exploded")

        monkeypatch.setattr("docx_cli.builder.validate.validate_document", boom)
        result = build_document(minimal_spec, str(out))
        assert not out.exists()
        assert result["ok"] is False
        assert result["status"] == "error"
