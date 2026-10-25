"""docx-cli v0.6.0 normalize vectors (v0.6 §5, audit F4).

F4: dry-run, backup, каталог с --yes, hyperlink-runs.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document

from docx_cli.fonts_normalize import normalize_document, normalize_path
from docx_cli.validate import validate_document


def _doc_with_calibri(tmp_path: Path, name: str = "c.docx") -> Path:
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Текст")
    r.font.name = "Calibri"
    r.font.size = __import__("docx.shared", fromlist=["Pt"]).Pt(12)
    path = tmp_path / name
    doc.save(str(path))
    return path


def _doc_with_hyperlink_calibri(tmp_path: Path, name: str = "hl.docx") -> Path:
    doc = Document()
    doc.add_paragraph()
    path = tmp_path / name
    doc.save(str(path))
    zin = zipfile.ZipFile(path)
    names = zin.namelist()
    docxml = zin.read("word/document.xml").decode("utf-8")
    docxml = docxml.replace(
        "</w:body>",
        '<w:p><w:hyperlink><w:r><w:rPr>'
        '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="24"/>'
        '</w:rPr><w:t>link</w:t></w:r></w:hyperlink></w:p></w:body>')
    out = tmp_path / ("hl_" + name)
    with zipfile.ZipFile(out, "w") as zout:
        for n in names:
            zout.writestr(n, docxml if n == "word/document.xml" else zin.read(n))
    zin.close()
    return out


def test_v060_f4_normalize_dry_run_no_write(tmp_path: Path) -> None:
    """F4: --dry-run показывает изменения, файл не трогает, .bak не создаёт."""
    f = _doc_with_calibri(tmp_path)
    original = f.read_bytes()
    res = normalize_document(str(f), dry_run=True)
    assert res["dry_run"] is True
    assert res["backup"] is None
    assert res["changes"] >= 1
    assert f.read_bytes() == original
    assert not (Path(str(f) + ".bak")).exists()


def test_v060_f4_normalize_backup_created(tmp_path: Path) -> None:
    """F4: по умолчанию при записи создаётся .bak рядом."""
    f = _doc_with_calibri(tmp_path)
    res = normalize_document(str(f))
    assert res["backup"] == str(f) + ".bak"
    assert Path(str(f) + ".bak").exists()
    assert "Calibri" not in f.read_text(encoding="utf-8", errors="ignore")


def test_v060_f4_normalize_no_backup(tmp_path: Path) -> None:
    """F4: --no-backup не создаёт .bak."""
    f = _doc_with_calibri(tmp_path)
    res = normalize_document(str(f), backup=False)
    assert res["backup"] is None
    assert not Path(str(f) + ".bak").exists()


def test_v060_f4_normalize_dir_requires_yes(tmp_path: Path) -> None:
    """F4: каталоговый режим требует --yes."""
    d = tmp_path / "docs"
    d.mkdir()
    doc = Document()
    doc.add_paragraph("x")
    (d / "a.docx").parent  # noop
    doc.save(str(d / "a.docx"))
    import pytest
    with pytest.raises(ValueError, match="invalid_usage"):
        normalize_path(str(d))
    res = normalize_path(str(d), yes=True)
    assert res["count"] == 1
    assert res["dry_run"] is False


def test_v060_f4_normalize_hyperlink_runs(tmp_path: Path) -> None:
    """F4: runs внутри w:hyperlink нормализуются."""
    f = _doc_with_hyperlink_calibri(tmp_path)
    assert "Calibri" in f.read_text(encoding="utf-8", errors="ignore")
    res = normalize_document(str(f))
    assert res["changes"] >= 1
    assert "Calibri" not in f.read_text(encoding="utf-8", errors="ignore")


def test_v060_f4_validate_hyperlink_runs(tmp_path: Path) -> None:
    """F4: валидатор проверяет runs внутри w:hyperlink."""
    f = _doc_with_hyperlink_calibri(tmp_path)
    res = validate_document(str(f))
    errors = res["errors"]
    assert any("font_name" in e and "Calibri" in e for e in errors)
