"""docx-cli v0.5.3 приёмка (внешний аудит, фиксы 1-5).

Фикстура — замороженный синтетический t.md из аудита,
вся анатомия дефектов: подзаголовок «## 1.» с утечкой маркера, вторая строка
титула в body, перечни «1) 2)», «ПРОШУ:» в body justify; плюс склейка
нумерации [1|1] в validate.

Фиксы: 1) парсер md_court; 2) инвариант md-build->validate ok:true;
3) счётчик перечней (рестарт [1..k] любой длины, границы ПРОШУ:/Приложения:
без учёта регистра/двоеточия); 4) normalize чинит красную строку корпуса;
5) validate --mode gost|review (геометрия в review — только warnings).
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from docx_cli.fonts_normalize import normalize_document
from docx_cli.md_court import md_build
from docx_cli.validate import _classify_paragraphs, validate_document

PYTHON = sys.executable
ROOT = Path(__file__).resolve().parents[1]

# --- замороженная фикстура: синтетический t.md из внешнего аудита ---
FIXTURE_MD = """ИСКОВОЕ ЗАЯВЛЕНИЕ
к апелляционной жалобе по делу № 123
Истец: Иванов И.И.

## 1. Нарушение ответчиком норм права
Согласно договору, ответчик нарушил сроки.

1) первое основание иска;
2) второе основание иска.

ПРОШУ:

1. Взыскать долг.
2. Взыскать неустойку.

Приложения:
1. Квитанция об уплате пошлины.

1. Опись документов.
"""


def build_fixture(tmp_path: Path) -> Path:
    md = tmp_path / "t.md"
    md.write_text(FIXTURE_MD, encoding="utf-8")
    out = tmp_path / "t.docx"
    md_build(str(md), str(out), "pytest")
    return out


def find_par(doc, needle: str):
    return next(p for p in doc.paragraphs if needle in p.text)


# --- фикс 1: парсер md_court --------------------------------------------------

def test_v053_subheading_marker_not_leaked(tmp_path):
    out = build_fixture(tmp_path)
    doc = docx.Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert not any("##" in t for t in texts), "маркер '##' утёк в текст"
    assert any("1. Нарушение ответчиком норм права" in t for t in texts)


def test_v053_subheading_bold_centered_no_fli(tmp_path):
    out = build_fixture(tmp_path)
    doc = docx.Document(str(out))
    p = find_par(doc, "Нарушение ответчиком норм права")
    assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p.paragraph_format.first_line_indent in (None, 0) or \
        abs(p.paragraph_format.first_line_indent.cm) < 0.05
    assert all(r.bold for r in p.runs if r.text)


def test_v053_title_second_line_stays_in_title_block(tmp_path):
    out = build_fixture(tmp_path)
    doc = docx.Document(str(out))
    p = find_par(doc, "к апелляционной жалобе по делу")
    assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    fl = p.paragraph_format.first_line_indent
    assert fl is None or int(fl) == 0
    roles = _classify_paragraphs(doc, None)
    idx = next(i for i, par in enumerate(doc.paragraphs)
               if "к апелляционной жалобе по делу" in par.text)
    assert roles.get(idx) == "title"


def test_v053_parenthesized_lists_hanging_geometry(tmp_path):
    out = build_fixture(tmp_path)
    doc = docx.Document(str(out))
    for needle in ("1) первое основание", "2) второе основание"):
        p = find_par(doc, needle)
        assert abs(p.paragraph_format.left_indent.cm - 1.25) <= 0.05
        assert abs(p.paragraph_format.first_line_indent.cm + 0.75) <= 0.05


def test_v053_proshu_centered_without_indent(tmp_path):
    out = build_fixture(tmp_path)
    doc = docx.Document(str(out))
    p = find_par(doc, "ПРОШУ:")
    assert p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(not (r.text and not r.bold) for r in p.runs)  # bold-маркер
    fl = p.paragraph_format.first_line_indent
    assert fl is None or int(fl) == 0


# --- фикс 2: инвариант самосогласованности ----------------------------------

def test_v053_invariant_md_build_then_validate_ok(tmp_path):
    out = build_fixture(tmp_path)
    res = validate_document(str(out))
    assert res["errors"] == [], res["errors"]


def test_v053_invariant_cli_md_build_validate_roundtrip(tmp_path):
    md = tmp_path / "t.md"
    md.write_text(FIXTURE_MD, encoding="utf-8")
    out = tmp_path / "t.docx"
    r1 = subprocess.run(
        [PYTHON, "-m", "docx_cli", "md-build", str(md), str(out),
         "--author", "pytest"],
        capture_output=True, text=True, cwd=str(ROOT))
    assert r1.returncode == 0, r1.stderr
    r2 = subprocess.run(
        [PYTHON, "-m", "docx_cli", "validate", str(out), "--mode", "gost"],
        capture_output=True, text=True, cwd=str(ROOT))
    env = json.loads(r2.stdout.splitlines()[0])
    assert r2.returncode == 0, env
    assert env["ok"] is True, env


# --- фикс 3: счётчик перечней ------------------------------------------------

def _mk_list_doc(tmp_path, lines) -> Path:
    d = docx.Document()
    for t in lines:
        d.add_paragraph(t)
    out = tmp_path / "seq.docx"
    d.save(str(out))
    return out


def test_v053_sequence_single_item_restart_ok(tmp_path):
    """Независимые одноэлементные перечни [1] [1] [1] — легитимны (в v0.5.2
    давали «out of order»: рестарт требовал предыдущую группу длиной >=2)."""
    out = _mk_list_doc(tmp_path, ["1. первый перечень", "1. второй перечень",
                                  "1. третий перечень"])
    res = validate_document(str(out))
    assert [e for e in res["errors"] if "numbered_list_sequence" in e] == []


def test_v053_sequence_proshu_boundary_case_insensitive(tmp_path):
    out = _mk_list_doc(tmp_path, ["1. основание", "2. основание",
                                  "прошу:", "1. взыскать"])
    res = validate_document(str(out))
    assert [e for e in res["errors"] if "numbered_list_sequence" in e] == []


def test_v053_sequence_attachments_boundary_no_colon(tmp_path):
    out = _mk_list_doc(tmp_path, ["1. взыскать", "приложения.", "1. копия"])
    res = validate_document(str(out))
    assert [e for e in res["errors"] if "numbered_list_sequence" in e] == []


def test_v053_sequence_true_gap_still_fails(tmp_path):
    """Fail-closed сохранён: разрыв [1,3,2] внутри группы — по-прежнему error."""
    out = _mk_list_doc(tmp_path, ["1. a", "3. c", "2. b"])
    res = validate_document(str(out))
    assert any("numbered_list_sequence" in e for e in res["errors"])


def test_v053_fixture_two_independent_lists_ok(tmp_path):
    out = build_fixture(tmp_path)
    doc = docx.Document(str(out))
    res = validate_document(str(out))
    assert [e for e in res["errors"] if "numbered_list_sequence" in e] == []
    _ = doc  # фикстура содержит [1|1] в блоке Приложения: без границы-маркера


# --- фикс 4: normalize красной строки ----------------------------------------

def test_v053_normalize_repairs_body_indent_keeps_lists(tmp_path):
    out = build_fixture(tmp_path)
    doc = docx.Document(str(out))
    body = find_par(doc, "Согласно договору")
    body.paragraph_format.first_line_indent = None
    other = find_par(doc, "Истец: Иванов")
    other.paragraph_format.first_line_indent = Cm(0.4)  # чужой отступ
    doc.save(str(out))
    res = normalize_document(str(out))
    assert res["first_line_indent_fixed"] == 2
    doc2 = docx.Document(str(out))
    assert abs(find_par(doc2, "Согласно договору")
               .paragraph_format.first_line_indent.cm - 1.25) <= 0.05
    assert abs(find_par(doc2, "Истец: Иванов")
               .paragraph_format.first_line_indent.cm - 1.25) <= 0.05
    # списки не тронуты (fli<0 — висячая геометрия)
    assert abs(find_par(doc2, "1) первое основание")
               .paragraph_format.first_line_indent.cm + 0.75) <= 0.05
    # титул/центр/подзаголовок не тронуты
    sub = find_par(doc2, "Нарушение ответчиком норм права")
    assert sub.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert sub.paragraph_format.first_line_indent in (None, 0) or \
        int(sub.paragraph_format.first_line_indent) == 0
    v = validate_document(str(out))
    assert v["errors"] == [], v["errors"]


def test_v053_normalize_keeps_header_before_title_untouched(tmp_path):
    d = docx.Document()
    hp = d.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.add_run("В райсуд, адрес: г. Москва")
    tp = d.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run("ХОДАТАЙСТВО")
    r.bold = True
    bp = d.add_paragraph()
    bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bp.paragraph_format.line_spacing = 1.5
    bp.add_run("Ходатайство о приобщении.")
    f = tmp_path / "hdr.docx"
    d.save(str(f))
    normalize_document(str(f))
    doc2 = docx.Document(str(f))
    assert doc2.paragraphs[0].paragraph_format.first_line_indent in (None, 0) \
        or int(doc2.paragraphs[0].paragraph_format.first_line_indent) == 0
    assert abs(doc2.paragraphs[2].paragraph_format.first_line_indent.cm
               - 1.25) <= 0.05


# --- фикс 5: режимы геометрии validate ---------------------------------------

def _foreign_doc(tmp_path: Path, top_cm: float = 3.5, placeholder: bool = False,
                 line_spacing: float = 1.15, indent=Pt(0)) -> Path:
    """Заведомо неГОСТовский входящий документ: поля за полосой review,
    шрифт 13pt (в полосе 10–14), красная строка отсутствует/чужая."""
    d = docx.Document()
    s = d.sections[0]
    s.top_margin = Cm(top_cm)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.0)
    s.right_margin = Cm(2.0)
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(13)
    tp = d.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run("ОТЗЫВ")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    bp = d.add_paragraph()
    bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bp.paragraph_format.line_spacing = line_spacing
    bp.paragraph_format.first_line_indent = indent
    text = "[УКАЗАТЬ сумму требований]." if placeholder else "Возражения против иска."
    rb = bp.add_run(text)
    rb.font.name = "Times New Roman"
    rb.font.size = Pt(13)
    out = tmp_path / "foreign.docx"
    d.save(str(out))
    return out


def test_v053_review_geometry_degrades_to_warnings(tmp_path):
    out = _foreign_doc(tmp_path, top_cm=3.5, indent=Cm(0.3))
    strict = validate_document(str(out))  # gost по умолчанию
    assert any("margin_top" in e for e in strict["errors"])
    res = validate_document(str(out), mode="review")
    assert res["errors"] == [], res["errors"]
    assert res["mode"] == "review"
    assert any("review_margin_top" in w for w in res["warnings"])
    assert any("review_first_line_indent" in w for w in res["warnings"])
    assert res["checks"]["margins"] == "warn"
    assert res["checks"]["fonts"] == "pass"  # 13pt в полосе 10–14


def test_v053_review_absent_indent_with_spacing_ok(tmp_path):
    """Красная строка отсутствует, но межпараграфный интервал явный ->
    предупреждения по отступу нет."""
    d = docx.Document()
    tp = d.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.add_run("ЗАЯВЛЕНИЕ")
    bp = d.add_paragraph()
    bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bp.paragraph_format.first_line_indent = None
    bp.paragraph_format.space_after = Pt(12)  # явный межпараграфный интервал
    bp.add_run("Абзац без красной строки, но с интервалом.")
    out = tmp_path / "spacing.docx"
    d.save(str(out))
    res = validate_document(str(out), mode="review")
    assert not any("first_line_indent" in w for w in res["warnings"]), res["warnings"]


def test_v053_review_placeholder_is_blocking_content_error(tmp_path):
    out = _foreign_doc(tmp_path, placeholder=True)
    res = validate_document(str(out), mode="review")
    assert any("placeholder_present" in e for e in res["errors"])


def test_v053_review_sequence_error_still_blocks(tmp_path):
    """Грубые структурные нарушения остаются errors и в review."""
    d = docx.Document()
    d.add_paragraph("1. a")
    d.add_paragraph("3. c")
    out = tmp_path / "bad_seq.docx"
    d.save(str(out))
    res = validate_document(str(out), mode="review")
    assert any("numbered_list_sequence" in e for e in res["errors"])


def test_v053_review_en_dash_still_blocks(tmp_path):
    d = docx.Document()
    d.add_paragraph("период 2020\u20132026 годы")
    out = tmp_path / "dash.docx"
    d.save(str(out))
    res = validate_document(str(out), mode="review")
    assert any("en_dash" in e for e in res["errors"])


def test_v053_gost_is_default_and_strict(tmp_path):
    out = _foreign_doc(tmp_path)
    default = validate_document(str(out))
    by_mode = validate_document(str(out), mode="gost")
    assert default["errors"] == by_mode["errors"]
    assert default["checks"]["margins"] in ("fail",)
    assert by_mode["mode"] == "gost"


def test_v053_unknown_mode_fails_closed(tmp_path):
    out = _foreign_doc(tmp_path)
    with pytest.raises(ValueError):
        validate_document(str(out), mode="nonsense")


def test_v053_cli_mode_flag(tmp_path):
    out = _foreign_doc(tmp_path, indent=Cm(0.3))
    r = subprocess.run(
        [PYTHON, "-m", "docx_cli", "validate", str(out), "--mode", "review"],
        capture_output=True, text=True, cwd=str(ROOT))
    env = json.loads(r.stdout.splitlines()[0])
    assert r.returncode == 0
    assert env["ok"] is True
    assert env["metadata"]["warnings"], "ожидались геометрические warnings"
    bad = _foreign_doc(tmp_path, placeholder=True)
    r2 = subprocess.run(
        [PYTHON, "-m", "docx_cli", "validate", str(bad), "--mode", "review"],
        capture_output=True, text=True, cwd=str(ROOT))
    env2 = json.loads(r2.stdout.splitlines()[0])
    assert r2.returncode == 1
    assert env2["ok"] is False
