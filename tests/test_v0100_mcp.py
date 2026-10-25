"""v0.1.0 §1 A-05: MCP stdio-сервер docx (JSON-RPC 2.0, паттерн sud).

8 тулов, permission read|write во всех схемах, write требует confirm=true
(needs_confirmation без сайд-эффекта), CLI `docx-cli mcp` / `mcp --tools`.
"""
from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

from docx_cli.mcp_server import (
    TOOLS,
    handle_request,
    serve_stdio,
    tool_list,
)

EXPECTED_TOOLS = {
    "docx_build", "docx_validate", "docx_scan", "docx_normalize",
    "docx_md_build", "docx_doctor", "docx_agent_read", "docx_verify_forbidden",
}
WRITE_TOOLS = {"docx_build", "docx_normalize", "docx_md_build"}


def _spec_dict(tmp_path: Path) -> dict:
    return {
        "profile": "gost",
        "output": {"file_name": "mcp-out.docx", "author": "Тест Т.Т."},
        "header": {"align": "right", "lines": ["В тестовый суд"]},
        "title": {"main": "ТЕСТОВЫЙ ДОКУМЕНТ"},
        "body": [{"type": "paragraph", "text": "Проверка MCP.", "inline_markdown": False}],
    }


def _call(name: str, arguments: dict) -> dict:
    resp = handle_request({"jsonrpc": "2.0", "id": 1, "method": "tools/call",
                           "params": {"name": name, "arguments": arguments}})
    assert resp["result"]["content"][0]["type"] == "text"
    return resp


def _payload(resp: dict) -> dict:
    return json.loads(resp["result"]["content"][0]["text"])


# ---------------------------------------------------------------- схемы


def test_v0100_mcp_tools_list_eight_with_schemas_and_permission():
    tl = tool_list()["tools"]
    assert {t["name"] for t in tl} == EXPECTED_TOOLS
    assert len(tl) == 8
    for t in tl:
        assert t["permission"] in ("read", "write")
        schema = t["inputSchema"]
        assert schema["type"] == "object"
        assert schema["additionalProperties"] is False
        # permission-поле объявлено в КАЖДОЙ схеме
        prop = schema["properties"]["permission"]
        assert set(prop["enum"]) == {"read", "write"}


def test_v0100_mcp_write_schemas_declare_confirm():
    for t in tool_list()["tools"]:
        if t["permission"] == "write":
            assert "confirm" in t["inputSchema"]["properties"]
            assert t["name"] in WRITE_TOOLS
        else:
            assert "confirm" not in t["inputSchema"]["properties"]


def test_v0100_mcp_permissions_split():
    read = {t["name"] for t in TOOLS if t["permission"] == "read"}
    write = {t["name"] for t in TOOLS if t["permission"] == "write"}
    assert write == WRITE_TOOLS
    assert read == EXPECTED_TOOLS - WRITE_TOOLS


# ---------------------------------------------------------------- confirm-гейт


def test_v0100_build_without_confirm_needs_confirmation_and_no_file(tmp_path):
    out = tmp_path / "never.docx"
    resp = _call("docx_build", {"spec": _spec_dict(tmp_path), "out": str(out)})
    assert resp["result"]["isError"] is False
    payload = _payload(resp)
    assert payload["status"] == "needs_confirmation"
    assert payload["permission"] == "write"
    assert not out.exists()


def test_v0100_build_with_confirm_writes_file(tmp_path):
    out = tmp_path / "yes.docx"
    resp = _call("docx_build", {"spec": _spec_dict(tmp_path), "out": str(out),
                                "confirm": True})
    payload = _payload(resp)
    assert payload["ok"] is True
    assert payload["status"] in ("ok", "degraded")
    assert out.exists()


def test_v0100_md_build_without_confirm_no_file(tmp_path):
    out = tmp_path / "md.docx"
    resp = _call("docx_md_build", {"md": "ХОДАТАЙСТВО\n\nПрошу.\n", "out": str(out)})
    assert _payload(resp)["status"] == "needs_confirmation"
    assert not out.exists()


def test_v0100_normalize_without_confirm_no_write(tmp_path):
    import docx
    p = tmp_path / "calibri.docx"
    d = docx.Document()
    d.add_paragraph("Calibri текст")
    d.save(str(p))
    mtimes = p.stat().st_mtime_ns
    resp = _call("docx_normalize", {"path": str(p)})
    assert _payload(resp)["status"] == "needs_confirmation"
    assert p.stat().st_mtime_ns == mtimes  # файл не тронут


# ---------------------------------------------------------------- read-тулы


def test_v0100_validate_tool_on_built_file(tmp_path):
    out = tmp_path / "v.docx"
    _call("docx_build", {"spec": _spec_dict(tmp_path), "out": str(out), "confirm": True})
    resp = _call("docx_validate", {"file": str(out)})
    payload = _payload(resp)
    assert payload["ok"] is True
    assert payload["status"] in ("ok", "degraded")


def test_v0100_scan_tool_clean_on_synthetic(tmp_path):
    import docx
    p = tmp_path / "clean.docx"
    d = docx.Document()
    d.add_paragraph("Чистый текст.")
    d.save(str(p))
    payload = _payload(_call("docx_scan", {"file": str(p)}))
    assert payload["clean"] is True


def test_v0100_agent_read_tool_window(tmp_path):
    out = tmp_path / "ar.docx"
    _call("docx_build", {"spec": _spec_dict(tmp_path), "out": str(out), "confirm": True})
    payload = _payload(_call("docx_agent_read", {"file": str(out), "start": 0,
                                                 "max_chars": 50}))
    assert payload["ok"] is True
    assert "blocks" in payload


def test_v0100_verify_forbidden_hit_and_miss(tmp_path):
    out = tmp_path / "vf.docx"
    _call("docx_build", {"spec": _spec_dict(tmp_path), "out": str(out), "confirm": True})
    miss = _payload(_call("docx_verify_forbidden", {"docx": str(out),
                                                    "forbid": ["нет-такой-строки"]}))
    assert miss["ok"] is True and miss["count"] == 0
    hit = _payload(_call("docx_verify_forbidden", {"docx": str(out),
                                                   "forbid": ["Проверка MCP."]}))
    assert hit["ok"] is False and hit["count"] >= 1


def test_v0100_doctor_tool_read(tmp_path):
    payload = _payload(_call("docx_doctor", {}))
    assert "checks" in payload["data"]


# ---------------------------------------------------------------- конверт/транспорт


def test_v0100_unknown_method_minus32601():
    resp = handle_request({"jsonrpc": "2.0", "id": 7, "method": "bogus/method"})
    assert resp["error"]["code"] == -32601


def test_v0100_malformed_line_returns_32700_and_loop_survives():
    stdin = io.StringIO('}{битая строка\n{"jsonrpc":"2.0","id":2,"method":"tools/list"}\n')
    stdout = io.StringIO()
    serve_stdio(stdin, stdout)
    lines = [l for l in stdout.getvalue().splitlines() if l.strip()]
    assert len(lines) == 2
    first = json.loads(lines[0])
    assert first["error"]["code"] == -32700
    second = json.loads(lines[1])
    assert len(second["result"]["tools"]) == 8  # цикл не умер


def test_v0100_initialize_serverinfo_version_tracks_package():
    import docx_cli
    resp = handle_request({"jsonrpc": "2.0", "id": 1, "method": "initialize",
                           "params": {}})
    assert resp["result"]["serverInfo"]["version"] == docx_cli.__version__


def test_v0100_unknown_argument_field_is_error_not_crash(tmp_path):
    resp = _call("docx_validate", {"file": "x.docx", "bogus": 1})
    assert resp["result"]["isError"] is True


def test_v0100_permission_field_mismatch_is_error(tmp_path):
    # объявленное read-право write-вызовом — явная ошибка, не молчание
    resp = _call("docx_build", {"spec": _spec_dict(tmp_path),
                                "out": str(tmp_path / "p.docx"),
                                "permission": "read", "confirm": True})
    assert resp["result"]["isError"] is True


def test_v0100_notification_initialized_returns_no_response():
    resp = handle_request({"jsonrpc": "2.0",
                           "method": "notifications/initialized", "params": {}})
    assert resp is None


# ---------------------------------------------------------------- CLI


def test_v0100_cli_mcp_tools_prints_valid_json_8_tools():
    import subprocess
    import sys
    cp = subprocess.run([sys.executable, "-m", "docx_cli", "mcp", "--tools"],
                        capture_output=True, text=True,
                        cwd=str(Path(__file__).parent.parent))
    assert cp.returncode == 0
    data = json.loads(cp.stdout)
    assert len(data["tools"]) == 8
    assert all("permission" in t for t in data["tools"])


def test_v0100_cli_mcp_stdio_reads_stdin():
    import subprocess
    import sys
    inp = ('{"jsonrpc":"2.0","id":1,"method":"initialize","params":{}}\n'
           '{"jsonrpc":"2.0","id":2,"method":"tools/list"}\n')
    cp = subprocess.run([sys.executable, "-m", "docx_cli", "mcp"],
                        input=inp, capture_output=True, text=True,
                        cwd=str(Path(__file__).parent.parent))
    assert cp.returncode == 0
    lines = [l for l in cp.stdout.splitlines() if l.strip()]
    assert len(lines) == 2
    assert json.loads(lines[0])["result"]["serverInfo"]["name"] == "docx-cli"
