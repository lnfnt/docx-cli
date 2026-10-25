"""docx-cli normalize — приведение шрифтов/интервалов к Times New Roman."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.run import Run

TARGET_FONT = "Times New Roman"
GOST_FIRST_LINE_INDENT_CM = 1.25
_TOL_CM = 0.05
_KNOWN_FONTS = {
    "calibri", "cambria", "arial", "verdana", "tahoma", "segoe ui",
    "helvetica", "liberation serif", "liberation sans", "dejavu sans",
    "pt serif", "pt sans",
}


def _all_runs(p):
    for r in p._p.iter(qn("w:r")):
        yield Run(r, p)


def _force_font(font_obj, size_pt: int | None = None) -> bool:
    changed = False
    name = (font_obj.name or "").lower()
    if name in _KNOWN_FONTS or (name and name != TARGET_FONT.lower()):
        font_obj.name = TARGET_FONT
        changed = True
    if size_pt is not None and font_obj.size is not None and font_obj.size != Pt(size_pt):
        font_obj.size = Pt(size_pt)
        changed = True
    return changed


def _repair_first_line_indents(doc, target_cm: float = GOST_FIRST_LINE_INDENT_CM) -> int:
    from docx_cli.validate import _classify_paragraphs
    roles = _classify_paragraphs(doc, None)
    fixed = 0
    for idx, p in enumerate(doc.paragraphs):
        if roles.get(idx) != "body":
            continue
        if not p.text.strip():
            continue
        pf = p.paragraph_format
        cur = pf.first_line_indent
        if cur is not None and abs(cur.cm - target_cm) <= _TOL_CM:
            continue
        pf.first_line_indent = Cm(target_cm)
        fixed += 1
    return fixed


def normalize_document(path, body_size: int = 12, header_size: int = 11,
                       dry_run: bool = False, backup: bool = True) -> dict:
    doc = Document(str(path))
    changed = 0

    if _force_font(doc.styles["Normal"].font, body_size):
        changed += 1
    for style in doc.styles:
        try:
            if style.font is not None and _force_font(style.font):
                changed += 1
        except Exception:  # noqa: BLE001
            continue
    for p in doc.paragraphs:
        if p.style and p.style.name and "Heading" in p.style.name:
            for run in _all_runs(p):
                changed += int(_force_font(run.font, header_size))
        else:
            for run in _all_runs(p):
                changed += int(_force_font(run.font, body_size))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in _all_runs(p):
                        changed += int(_force_font(run.font, body_size))
    for section in doc.sections:
        for container in (section.header, section.footer):
            for p in container.paragraphs:
                for run in _all_runs(p):
                    changed += int(_force_font(run.font, header_size))

    indent_fixed = _repair_first_line_indents(doc)
    changed += indent_fixed

    backup_path = None
    if not dry_run:
        if backup:
            backup_path = str(path) + ".bak"
            shutil.copy2(str(path), backup_path)
        doc.save(str(path))
    return {"file": str(path), "changes": changed,
            "first_line_indent_fixed": indent_fixed,
            "dry_run": dry_run, "backup": backup_path}


def normalize_path(path, body_size: int = 12, header_size: int = 11,
                   dry_run: bool = False, backup: bool = True,
                   yes: bool = False) -> dict:
    p = Path(path)
    if p.is_file():
        res = normalize_document(p, body_size, header_size, dry_run=dry_run, backup=backup)
        return {"normalized": [res], "count": 1, "dry_run": dry_run}
    if p.is_dir():
        if not yes:
            raise ValueError("invalid_usage: directory mode requires --yes")
        files = [f for f in sorted(p.glob("*.docx")) if not f.name.startswith("~$")]
        if not files:
            raise FileNotFoundError(f"нет .docx в каталоге: {p}")
        results = [normalize_document(f, body_size, header_size, dry_run=dry_run, backup=backup)
                   for f in files]
        return {"normalized": results, "count": len(results), "dry_run": dry_run}
    raise FileNotFoundError(f"not found: {p}")
