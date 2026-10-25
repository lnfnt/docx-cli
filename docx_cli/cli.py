"""CLI entry point for docx-cli."""
from __future__ import annotations

import argparse
import json
import os
import platform
import sys
from pathlib import Path
from typing import Any

from docx_cli import validate
from docx_cli.builder import build_document
from docx_cli.errors import DocxCreatorError
from docx_cli.spec import DocSpec, now_iso


def _envelope(
    ok: bool,
    status: str,
    reason: str | None,
    data: dict[str, Any] | None = None,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "ok": ok,
        "status": status,
        "reason": reason,
        "parsed_at": now_iso(),
        "data": data or {},
        "metadata": metadata or {},
    }


def _load_spec(path: str) -> DocSpec:
    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)
    return DocSpec.from_dict(raw)


def cmd_build(args: argparse.Namespace) -> int:
    try:
        spec = _load_spec(args.spec)
        if args.profile:
            spec.profile = args.profile
    except DocxCreatorError as exc:
        env = _envelope(False, "error", f"spec_error: {exc}")
        print(json.dumps(env, ensure_ascii=False))
        return 1
    except Exception as exc:
        env = _envelope(False, "error", f"spec_load_failed: {exc}")
        print(json.dumps(env, ensure_ascii=False))
        return 1

    try:
        result = build_document(spec, args.output)
    except DocxCreatorError as exc:
        env = _envelope(False, "error", str(exc))
        print(json.dumps(env, ensure_ascii=False))
        return 1
    except Exception as exc:
        env = _envelope(False, "error", f"build_failed: {exc}")
        print(json.dumps(env, ensure_ascii=False))
        return 1

    print(json.dumps(result, ensure_ascii=False))
    if result["status"] == "ok":
        return 0
    if result["status"] == "degraded":
        return 2
    return 1


def cmd_profile_extract(args: argparse.Namespace) -> int:
    import json as _json
    from .profiles import ProfileError, extract
    try:
        prof = extract(args.sample)
    except (ProfileError, Exception) as exc:  # noqa: BLE001
        print(_json.dumps({"ok": False, "status": "error", "reason": str(exc)}, ensure_ascii=False)); return 1
    try:
        import pathlib
        pathlib.Path(args.out_json).write_text(
            _json.dumps(prof, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:  # noqa: BLE001
        print(_json.dumps({"ok": False, "status": "io_error", "reason": str(exc)}, ensure_ascii=False)); return 1
    print(_json.dumps({"ok": True, "status": "ok", "profile": args.out_json}, ensure_ascii=False))
    return 0


def cmd_profile_validate(args: argparse.Namespace) -> int:
    import json as _json
    from .profiles import ProfileError, load_profile, validate
    try:
        prof = load_profile(args.profile_json)
        problems = validate(args.file, prof)
    except (ProfileError, Exception) as exc:  # noqa: BLE001
        print(_json.dumps({"ok": False, "status": "error", "reason": str(exc)}, ensure_ascii=False)); return 1
    print(_json.dumps({"ok": len(problems) == 0, "status": "ok" if not problems else "problems",
                       "count": len(problems), "problems": problems}, ensure_ascii=False))
    return 0 if not problems else 1


def cmd_profile_repair(args: argparse.Namespace) -> int:
    import json as _json
    from .profiles import ProfileError, load_profile, repair
    try:
        prof = load_profile(args.profile_json)
        res = repair(args.file, prof, args.out)
    except Exception as exc:  # noqa: BLE001
        print(_json.dumps({"ok": False, "status": "error", "reason": str(exc)}, ensure_ascii=False)); return 1
    res.update({"ok": True, "status": "ok"})
    print(_json.dumps(res, ensure_ascii=False))
    return 0


def cmd_agent_read(args: argparse.Namespace) -> int:
    import json as _json
    from .agent_reader import AgentReadError, read
    try:
        res = read(args.file, args.start, args.max_chars)
    except AgentReadError as exc:
        msg = str(exc)
        status = "invalid_usage" if "invalid_usage" in msg else "error"
        print(_json.dumps({"ok": False, "status": status, "reason": msg}, ensure_ascii=False))
        return 2 if status == "invalid_usage" else 1
    res.update({"ok": True, "status": "ok"})
    print(_json.dumps(res, ensure_ascii=False))
    return 0


def cmd_verify_forbidden(args: argparse.Namespace) -> int:
    """desync-guard юриста: запрещённые строки не должны попасть в итоговый docx/md."""
    import json as _json
    import zipfile as _zf
    from pathlib import Path as _P
    from xml.etree import ElementTree as _ET
    from docx import Document as _D

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    # v0.5.1: --forbid теперь nargs='+' (группы значений); флэттим, сохраняя
    # совместимость со старые вызовами (список строк).
    patterns: list[str] = []
    for group in args.forbid or []:
        if isinstance(group, list):
            patterns.extend(group)
        else:
            patterns.append(group)
    if args.forbid_file:
        try:
            for line in _P(args.forbid_file).read_text(encoding="utf-8-sig").splitlines():
                line = line.strip()
                if line and not line.startswith("#"):
                    patterns.append(line)
        except OSError as exc:
            print(_json.dumps({"ok": False, "status": "io_error", "reason": "io_error",
                               "error": "io_error", "findings": [{"type": "io_error", "path": str(args.forbid_file), "reason": str(exc)}],
                               "count": 0, "hits": []}, ensure_ascii=False))
            return 1
    if not patterns:
        print(_json.dumps({"ok": False, "status": "invalid_usage", "reason": "no --forbid patterns"}, ensure_ascii=False))
        return 2

    # collect texts (same parts as scan: doc + headers/footers + footnotes/endnotes + tables)
    if not _P(args.docx).is_file():
        _emit_vf_error("io_error", str(args.docx), "file not found")
        return 1
    try:
        d = _D(str(args.docx))
    except Exception as exc:  # noqa: BLE001
        msg = str(exc).lower()
        if "zip" in msg or "not a zip" in msg or "no such file" in msg or "file not found" in msg:
            _emit_vf_error("io_error", str(args.docx), str(exc))
            return 1
        _emit_vf_error("xml_parse_error", str(args.docx), str(exc))
        return 1

    texts = []
    try:
        texts += [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        texts.append(cell.text)
        for section in d.sections:
            for part in (section.header, section.footer,
                         getattr(section, "first_page_header", None),
                         getattr(section, "first_page_footer", None),
                         getattr(section, "even_page_header", None),
                         getattr(section, "even_page_footer", None)):
                if part is None:
                    continue
                for p in part.paragraphs:
                    if p.text.strip():
                        texts.append(p.text)
                for t in part.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                texts.append(cell.text)
        # footnotes / endnotes from zip
        with _zf.ZipFile(str(args.docx)) as zf:
            for partname in ("word/footnotes.xml", "word/endnotes.xml"):
                if partname not in zf.namelist():
                    continue
                try:
                    root = _ET.fromstring(zf.read(partname))
                except _ET.ParseError as exc:
                    _emit_vf_error("xml_parse_error", partname, str(exc))
                    return 1
                for p in root.iter(f"{{{W}}}p"):
                    txt = "".join(t.text or "" for t in p.iter(f"{{{W}}}t"))
                    if txt.strip():
                        texts.append(txt)
    except FileNotFoundError as exc:
        _emit_vf_error("io_error", str(args.docx), str(exc))
        return 1
    except _zf.BadZipFile as exc:
        _emit_vf_error("io_error", str(args.docx), str(exc))
        return 1
    except OSError as exc:
        _emit_vf_error("io_error", str(args.docx), str(exc))
        return 1
    except _ET.ParseError as exc:
        _emit_vf_error("xml_parse_error", str(args.docx), str(exc))
        return 1

    if args.md and _P(args.md).is_file():
        try:
            texts += [l for l in _P(args.md).read_text(encoding="utf-8-sig").splitlines() if l.strip()]
        except OSError as exc:
            _emit_vf_error("io_error", str(args.md), str(exc))
            return 1
    elif args.md and _P(args.md).exists() is False and args.md:
        # md optional; ignore missing
        pass

    hits = []
    for pat in patterns:
        for txt in texts:
            if pat in txt:
                hits.append({"pattern": pat, "found_in": txt[:80]})
    print(_json.dumps({"ok": not hits, "status": "ok" if not hits else "forbidden_found",
                       "count": len(hits), "hits": hits[:20]}, ensure_ascii=False))
    return 0 if not hits else 1


def _emit_vf_error(kind: str, path: str, reason: str) -> None:
    import json as _json
    print(_json.dumps({"ok": False, "status": kind, "reason": kind,
                       "error": kind, "findings": [{"type": kind, "path": path, "reason": reason}],
                       "count": 0, "hits": []}, ensure_ascii=False))


def cmd_md_build(args: argparse.Namespace) -> int:
    """v0.4.0: Markdown -> процессуальный docx (собственные эвристики)."""
    import json as _json
    from .md_court import MdBuildError, md_build
    try:
        res = md_build(args.input_md, args.output, args.author)
    except (MdBuildError, Exception) as exc:  # noqa: BLE001
        print(_json.dumps({"ok": False, "status": "error", "reason": str(exc)},
                          ensure_ascii=False))
        return 1
    res.update({"ok": True, "status": "ok"})
    print(_json.dumps(res, ensure_ascii=False))
    return 0


def cmd_normalize(args: argparse.Namespace) -> int:
    """v0.4.0: принудительный Times New Roman + интервалы (собственные эвристики)."""
    import json as _json
    from pathlib import Path
    from .fonts_normalize import normalize_path

    p = Path(args.path)
    if p.is_dir() and not args.yes:
        print(_json.dumps({"ok": False, "status": "invalid_usage",
                           "reason": "directory mode requires --yes (pass --yes to process all .docx)"},
                          ensure_ascii=False))
        return 2
    try:
        res = normalize_path(args.path, body_size=args.body_size,
                             header_size=args.header_size, dry_run=args.dry_run,
                             backup=args.backup, yes=args.yes)
    except Exception as exc:  # noqa: BLE001 — fail-closed envelope
        print(_json.dumps({"ok": False, "status": "error", "reason": str(exc)},
                          ensure_ascii=False))
        return 1
    res.update({"ok": True, "status": "ok"})
    print(_json.dumps(res, ensure_ascii=False))
    return 0


def cmd_scan(args: argparse.Namespace) -> int:
    """v0.3.0: скрытые векторы атак в входящих docx (собственные эвристики).
    clean -> ok/exit 0; находки или битый файл -> error/exit 1 (fail-closed)."""
    import json as _json
    import shlex
    from pathlib import Path as _P
    from .inject_scan import scan_all

    try:
        res = scan_all(str(args.file))
    except Exception as exc:  # noqa: BLE001 — fail-closed envelope
        print(_json.dumps({"ok": False, "status": "error", "reason": str(exc)},
                          ensure_ascii=False))
        return 1
    if getattr(args, "a2a_args", False):
        count = len(res.get("findings", [])) if isinstance(res, dict) else 0
        file_path = str(args.file)
        scan_json_path = str(_P(file_path).with_suffix(".scan.json"))
        verify_cmd = f"docx-cli scan {shlex.quote(file_path)}"
        a2a_command = (
            f'a2a log "scan findings: {file_path} ({count})" '
            f'--status done --artifacts {shlex.quote(scan_json_path)} '
            f'--verify {shlex.quote(verify_cmd)} --verify-run'
        )
        res["a2a_command"] = a2a_command
        res.setdefault("data", {})["a2a_command"] = a2a_command
    print(_json.dumps(res, ensure_ascii=False))
    return 0 if res["clean"] else 1


def cmd_validate(args: argparse.Namespace) -> int:
    path = args.file
    if not os.path.exists(path):
        env = _envelope(False, "io_error", f"io_error: file_not_found: {path}")
        print(json.dumps(env, ensure_ascii=False))
        return 1

    try:
        validation = validate.validate_document(path, mode=args.mode)
    except Exception as exc:
        env = _envelope(False, "error", f"validate_failed: {exc}")
        print(json.dumps(env, ensure_ascii=False))
        return 1

    ok = not validation["errors"]
    has_warnings = bool(validation["warnings"])
    if ok and has_warnings:
        status = "degraded"
        reason = "warnings_present"
    elif ok:
        status = "ok"
        reason = None
    else:
        status = "error"
        reason = "validation_failed: " + "; ".join(validation["errors"])

    env = _envelope(
        ok,
        status,
        reason,
        data={"file": path},
        metadata=validation,
    )
    print(json.dumps(env, ensure_ascii=False))

    if not ok:
        return 1
    if args.strict and has_warnings:
        return 2
    if args.strict and validation.get("skipped"):
        return 2
    return 0


def cmd_doctor(args: argparse.Namespace) -> int:
    import shutil as _shutil
    import subprocess as _sp

    checks: dict[str, Any] = {}
    errors: list[str] = []
    warnings: list[str] = []

    checks["python_version"] = platform.python_version()
    if sys.version_info < (3, 11):
        errors.append("python_version: requires >=3.11")

    try:
        import docx
        checks["python_docx"] = docx.__version__
    except Exception as exc:
        errors.append(f"python_docx_import_failed: {exc}")

    try:
        import lxml
        checks["lxml"] = lxml.__version__
    except Exception as exc:
        errors.append(f"lxml_import_failed: {exc}")

    try:
        import pydantic
        checks["pydantic"] = pydantic.VERSION
    except Exception as exc:
        errors.append(f"pydantic_import_failed: {exc}")

    try:
        import docx_cli
        checks["docx_cli"] = getattr(docx_cli, "__version__", "unknown")
    except Exception as exc:
        errors.append(f"docx_cli_import_failed: {exc}")

    cwd = Path.cwd()
    if os.access(cwd, os.W_OK):
        checks["write_permissions"] = str(cwd)
    else:
        errors.append(f"write_permissions: no write access to {cwd}")

    verify_sig = bool(getattr(args, "verify_signature", False))
    dist = Path(getattr(args, "dist", "dist") or "dist")
    signature_status = None
    sig_reason = None
    if verify_sig:
        sums_path = dist / "SHA256SUMS"
        sig_path = None
        for cand in (dist / "SHA256SUMS.asc", dist / "SHA256SUMS.sig", dist / "SHA256SUMS.gpg"):
            if cand.is_file():
                sig_path = cand
                break
        if _shutil.which("gpg") is None:
            signature_status = "unsigned"
            sig_reason = "gpg binary not found"
        elif not sums_path.is_file() or sig_path is None:
            signature_status = "unsigned"
            sig_reason = f"SHA256SUMS or paired .asc not found in {dist}"
        else:
            try:
                cp = _sp.run(["gpg", "--verify", str(sig_path), str(sums_path)],
                             capture_output=True, text=True)
            except Exception as exc:  # noqa: BLE001
                signature_status = "unsigned"
                sig_reason = f"gpg invocation failed: {exc}"
            else:
                if cp.returncode == 0:
                    # подпись валидна — ОБЯЗАТЕЛЬНО сверяем хэши перечисленных файлов
                    # (иначе подмена dist-артефакта пройдёт как ok: fail-open)
                    import hashlib
                    mismatch = []
                    for line in sums_path.read_text().splitlines():
                        parts = line.split(None, 1)
                        if len(parts) != 2:
                            continue
                        want, name = parts[0].strip(), parts[1].strip().lstrip("*")
                        f = dist / name
                        if not f.exists():
                            mismatch.append(name)
                        elif hashlib.sha256(f.read_bytes()).hexdigest() != want:
                            mismatch.append(name)
                    if mismatch:
                        signature_status = "fail"
                        sig_reason = f"sha256 mismatch for: {', '.join(mismatch[:5])}"
                    else:
                        signature_status = "ok"
                else:
                    signature_status = "fail"
                    sig_reason = (cp.stderr or cp.stdout or "gpg verify failed").strip()

        checks["signature"] = signature_status
        checks["verify_signature"] = signature_status
        checks["signature_status"] = signature_status

    ok = not errors
    if verify_sig:
        ok = ok and signature_status == "ok"

    if errors:
        status = "error"
    elif verify_sig:
        status = signature_status
    elif warnings:
        status = "degraded"
    else:
        status = "ok"

    reasons = list(errors)
    if verify_sig and sig_reason:
        reasons.append(sig_reason)
    reason = "; ".join(reasons) if reasons else None

    data = {"checks": checks}
    metadata = {"warnings": warnings, "errors": errors}
    if verify_sig:
        data["signature_status"] = signature_status
        data["verify_signature"] = signature_status
        data["signature"] = signature_status
        data["dist"] = str(dist)
        metadata["signature_status"] = signature_status
        metadata["verify_signature"] = signature_status
        metadata["signature"] = signature_status

    env = _envelope(
        ok,
        status,
        reason,
        data=data,
        metadata=metadata,
    )
    print(json.dumps(env, ensure_ascii=False))
    return 0 if ok else 1


def cmd_mcp(args: argparse.Namespace) -> int:
    """v0.1.0 A-05: docx-cli mcp — MCP stdio-сервер; --tools — tools/list dump."""
    from . import mcp_server
    if getattr(args, "tools", False):
        return mcp_server.run_tools_dump()
    return mcp_server.run_stdio_server()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="docx-cli",
        description="Генератор юридических документов .docx по машинному контракту.",
    )
    sub = parser.add_subparsers(dest="command", required=True)

    build_p = sub.add_parser("build", help="Собрать .docx из spec.json")
    build_p.add_argument("spec", help="Путь к spec.json")
    build_p.add_argument("output", help="Путь к output.docx")
    build_p.add_argument("--profile",
                         choices=["gost", "compact", "court", "vs_sud", "soyu_rayon"],
                         help="Переопределить profile в spec")
    build_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    build_p.set_defaults(func=cmd_build)

    val_p = sub.add_parser("validate", help="Проверить существующий .docx")
    val_p.add_argument("file", help="Путь к .docx")
    val_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    val_p.add_argument("--strict", action="store_true", help="Предупреждения -> exit 2")
    # v0.5.3 (фикс 5; согласовано с внешним рецензентом):
    # gost — строгая геометрия для собственных документов (по умолчанию,
    # прежнее поведение); review — входящие акты судов/оппонентов: геометрия
    # (поля 1.0–3.0 см, шрифт 10–14 pt, межстрочник 1.0–1.5, красная строка
    # 0.5–1.5 см либо отсутствие при явном межпараграфном интервале) — только
    # warnings; errors строго контентные.
    val_p.add_argument("--mode", choices=["gost", "review"], default="gost",
                       help="gost (default) — строгая геометрия; review — "
                            "входящие документы: геометрия только warnings")
    val_p.set_defaults(func=cmd_validate)

    doc_p = sub.add_parser("doctor", help="Проверить окружение")
    doc_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    doc_p.add_argument("--verify-signature", dest="verify_signature", action="store_true",
                       help="Проверить подпись SHA256SUMS в dist (--dist DIR)")
    doc_p.add_argument("--dist", default="dist", help="Каталог с SHA256SUMS и .asc (по умолчанию dist/)")
    doc_p.set_defaults(func=cmd_doctor)

    scan_p = sub.add_parser(
        "scan", help="Аудит входящего .docx на скрытые LLM-атаки "
        "(vanish/микрошрифт/белый текст/tracked changes/comments). "
        "Прогонять ДО передачи файла третьих лиц агенту.")
    scan_p.add_argument("file", help="Путь к .docx")
    scan_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    scan_p.add_argument("--a2a-args", dest="a2a_args", action="store_true",
                        help="Добавить a2a_command (bus-free) в конверт")
    scan_p.set_defaults(func=cmd_scan)

    mdb_p = sub.add_parser(
        "md-build", help="Собрать процессуальный .docx из Markdown "
        "(шапка-реквизиты, титул ИСК/ОТЗЫВ/ХОДАТАЙСТВО, **жирный**, списки).")
    mdb_p.add_argument("input_md", help="Путь к исходному .md")
    mdb_p.add_argument("output", help="Путь к output.docx")
    mdb_p.add_argument("--author", default="docx-cli")
    mdb_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    mdb_p.set_defaults(func=cmd_md_build)

    norm_p = sub.add_parser(
        "normalize", help="Нормализация шрифтов/интервалов .docx к Times New Roman "
        "(вычищает Calibri/Cambria/Arial). Каталог или файл.")
    norm_p.add_argument("path", help="Файл или каталог .docx")
    norm_p.add_argument("--body-size", type=int, default=12)
    norm_p.add_argument("--header-size", type=int, default=11)
    norm_p.add_argument("--dry-run", action="store_true",
                        help="Показать, что будет изменено, без записи файлов")
    norm_p.add_argument("--backup", dest="backup", action="store_true", default=True,
                        help="Создать .bak рядом при записи (по умолчанию)")
    norm_p.add_argument("--no-backup", dest="backup", action="store_false",
                        help="Не создавать .bak")
    norm_p.add_argument("--yes", action="store_true",
                        help="Подтвердить обработку каталога")
    norm_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    norm_p.set_defaults(func=cmd_normalize)

    pe_p = sub.add_parser("profile-extract", help="Снять профиль оформления с образца .docx")
    pe_p.add_argument("sample")
    pe_p.add_argument("out_json")
    pe_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    pe_p.set_defaults(func=cmd_profile_extract)

    pv_p = sub.add_parser("profile-validate", help="Проверить .docx по профилю оформления")
    pv_p.add_argument("file")
    pv_p.add_argument("profile_json")
    pv_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    pv_p.set_defaults(func=cmd_profile_validate)

    pr_p = sub.add_parser("profile-repair", help="Починить .docx в профиль -> новый файл (before/after)")
    pr_p.add_argument("file")
    pr_p.add_argument("profile_json")
    pr_p.add_argument("out")
    pr_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    pr_p.set_defaults(func=cmd_profile_repair)

    ar_p = sub.add_parser("agent-read", help="Ограниченное JSON-чтение docx (окно + курсор)")
    ar_p.add_argument("file")
    ar_p.add_argument("--start", type=int, default=0)
    ar_p.add_argument("--max", "--max-chars", dest="max_chars", type=int, default=4000)
    ar_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    ar_p.set_defaults(func=cmd_agent_read)

    vf_p = sub.add_parser("verify-forbidden", help="Проверить docx (+md) на запрещённые строки (desync-guard)")
    vf_p.add_argument("docx")
    vf_p.add_argument("md", nargs="?")
    vf_p.add_argument("--forbid", action="append", nargs="+", default=[],
                      metavar="СТРОКА")  # v0.5.1: --forbid A B C за один раз
    vf_p.add_argument("--forbid-file", dest="forbid_file")
    vf_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    vf_p.set_defaults(func=cmd_verify_forbidden)

    mcp_p = sub.add_parser(
        "mcp", help="Запустить MCP stdio-сервер (JSON-RPC 2.0). "
        "--tools — печатать tools/list (скриптуемый smoke) и выйти.")
    mcp_p.add_argument("--tools", action="store_true",
                       help="напечатать tools/list JSON и выйти")
    mcp_p.add_argument("--json", action="store_true", help="Вывести JSON (по умолчанию)")
    mcp_p.set_defaults(func=cmd_mcp)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
