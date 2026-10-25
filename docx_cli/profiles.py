"""docx-cli profile-*. — профили оформления: снять с образца / проверить / починить.

v0.5.0 (clean-room: собственный профиль-цикл — снять профиль с образца,
проверить, починить). Контракт: наш JSON-конверт, честные exit-коды.

Профиль — JSON с параметрами страницы и тела:
  {"page": {"w_cm":21,"h_cm":29.7}, "margins": {"left":3,"right":1.5,"top":2,"bottom":2},
   "font": {"name":"Times New Roman","size":12},
   "body": {"first_line_indent_cm":1.25,"line_spacing":1.5}}
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
import zipfile
from xml.etree import ElementTree as ET
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

# допуск сравнения в см (вещественные EMU шумят на 4-м знаке)
TOL_CM = 0.05
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _all_paragraphs(doc):
    out = []
    def add_table(t):
        for row in t.rows:
            for cell in row.cells:
                out.extend(cell.paragraphs)
                for nt in cell.tables:
                    add_table(nt)
    out.extend(doc.paragraphs)
    for t in doc.tables:
        add_table(t)
    for section in doc.sections:
        for part in (section.header, section.footer,
                     getattr(section, "first_page_header", None),
                     getattr(section, "first_page_footer", None),
                     getattr(section, "even_page_header", None),
                     getattr(section, "even_page_footer", None)):
            if part is None:
                continue
            try:
                out.extend(part.paragraphs)
                for t in part.tables:
                    add_table(t)
            except Exception:  # noqa: BLE001
                pass
    return out


def _labeled_paragraphs(doc):
    out = []
    for i, p in enumerate(doc.paragraphs):
        out.append((f"paragraph {i}", p))
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    out.append((f"table {ti} cell {ci} paragraph {pi}", p))
    for si, section in enumerate(doc.sections):
        for name, part in (("header", section.header), ("footer", section.footer)):
            for pi, p in enumerate(part.paragraphs):
                out.append((f"{name} paragraph {pi}", p))
            for ti, t in enumerate(part.tables):
                for ri, row in enumerate(t.rows):
                    for ci, cell in enumerate(row.cells):
                        for pi, p in enumerate(cell.paragraphs):
                            out.append((f"{name} table {ti} cell {ci} paragraph {pi}", p))
    return out


class ProfileError(Exception):
    """Fail-closed ошибка профиля."""


def load_profile(path: str | Path) -> dict:
    p = Path(path)
    if not p.is_file():
        raise ProfileError(f"profile not found: {p}")
    data = json.loads(p.read_text(encoding="utf-8-sig"))
    for key in ("margins", "font"):
        if key not in data:
            raise ProfileError(f"profile missing key: {key}")
    return data


def extract(sample_path: str | Path) -> dict:
    """Снять профиль оформления с образца .docx."""
    d = Document(str(sample_path))
    s = d.sections[0]
    normal = d.styles["Normal"]
    body_indent = None
    body_spacing = None
    for p in d.paragraphs:
        pf = p.paragraph_format
        if p.text.strip() and pf.first_line_indent and pf.first_line_indent > 0:
            body_indent = pf.first_line_indent.cm
            body_spacing = pf.line_spacing
            break

    # шрифт — только факт из rPr/styles (без фолбэка TNR)
    name = normal.font.name
    size = normal.font.size.pt if normal.font.size else None
    if name is None or size is None:
        for p in _all_paragraphs(d):
            for run in p.runs:
                if name is None and run.font.name:
                    name = run.font.name
                if size is None and run.font.size:
                    size = run.font.size.pt
    if name is None or size is None:
        try:
            with zipfile.ZipFile(str(sample_path)) as zf:
                for n in sorted(zf.namelist()):
                    if not (n.startswith("word/") and n.endswith(".xml")):
                        continue
                    base = n.split("/")[-1]
                    if base in ("styles.xml", "numbering.xml", "fontTable.xml", "settings.xml", "webSettings.xml"):
                        continue
                    if not (base == "document.xml" or base.startswith("header") or base.startswith("footer")
                            or base.startswith("footnote") or base.startswith("endnote")):
                        continue
                    root = ET.fromstring(zf.read(n))
                    for rpr in root.iter(f"{{{W}}}rPr"):
                        if name is None:
                            rf = rpr.find(f"{{{W}}}rFonts")
                            if rf is not None:
                                for attr in ("ascii", "hAnsi"):
                                    v = rf.get(f"{{{W}}}{attr}")
                                    if v:
                                        name = v
                                        break
                        if size is None:
                            sz = rpr.find(f"{{{W}}}sz")
                            if sz is not None:
                                val = sz.get(f"{{{W}}}val")
                                if val and val.isdigit():
                                    size = int(val) / 2
                        if name is not None and size is not None:
                            break
                    if name is not None and size is not None:
                        break
        except Exception:  # noqa: BLE001
            pass
    guessed = name is None
    if size is None:
        size = 12
    prof = {
        "page": {"w_cm": round(s.page_width.cm, 3), "h_cm": round(s.page_height.cm, 3)},
        "margins": {"left": round(s.left_margin.cm, 3), "right": round(s.right_margin.cm, 3),
                    "top": round(s.top_margin.cm, 3), "bottom": round(s.bottom_margin.cm, 3)},
        "font": {"name": name, "size": size, "guessed": guessed},
        "body": {"first_line_indent_cm": round(body_indent, 3) if body_indent else None,
                 "line_spacing": body_spacing if isinstance(body_spacing, float) else None},
        "guessed_font": guessed,
        "font_guessed": guessed,
    }
    return prof


def _cm(v) -> float | None:
    return round(v.cm, 3) if v is not None else None


def validate(docx_path: str | Path, profile: dict) -> list[dict]:
    """Список проблем (не чинит). Проблема: {where, param, got, want}."""
    d = Document(str(docx_path))
    problems: list[dict] = []
    s = d.sections[0]
    m, pg = profile["margins"], profile.get("page", {})
    checks = [
        ("section", "margin_left", _cm(s.left_margin), m["left"]),
        ("section", "margin_right", _cm(s.right_margin), m["right"]),
        ("section", "margin_top", _cm(s.top_margin), m["top"]),
        ("section", "margin_bottom", _cm(s.bottom_margin), m["bottom"]),
    ]
    if pg.get("w_cm"):
        checks.append(("section", "page_width", _cm(s.page_width), pg["w_cm"]))
    if pg.get("h_cm"):
        checks.append(("section", "page_height", _cm(s.page_height), pg["h_cm"]))
    f = d.styles["Normal"].font
    checks.append(("styles/Normal", "font_name", f.name, profile["font"]["name"]))
    if f.size is not None:
        checks.append(("styles/Normal", "font_size", f.size.pt, profile["font"]["size"]))
    for where, param, got, want in checks:
        if got is None or want is None:
            continue
        if isinstance(got, str):
            if got.lower() != str(want).lower():
                problems.append({"where": where, "param": param, "got": got, "want": want})
        elif abs(got - float(want)) > TOL_CM:
            problems.append({"where": where, "param": param, "got": got, "want": want})
    # run-level шрифты: если задан явный шрифт, не совпадающий с профилем — проблема
    prof_font_name = profile["font"].get("name")
    if prof_font_name:
        want_name = str(prof_font_name).lower()
        for where, p in _labeled_paragraphs(d):
            for run in p.runs:
                if run.font.name and run.font.name.lower() != want_name:
                    problems.append({"where": where, "param": "run_font_name",
                                     "got": run.font.name, "want": prof_font_name})
    body = profile.get("body", {})
    want_ind = body.get("first_line_indent_cm")
    want_sp = body.get("line_spacing")
    for i, p in enumerate(d.paragraphs):
        pf = p.paragraph_format
        if not p.text.strip():
            continue
        # body-абзацы: не трогаем проверкой заголовки (bold/центр) и шапку (явный левый отступ)
        runs_bold = any(r.bold for r in p.runs if r.bold)
        is_center = p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        if runs_bold or is_center or (pf.left_indent and pf.left_indent > 0):
            continue
        if want_ind and pf.first_line_indent and abs(pf.first_line_indent.cm - want_ind) > TOL_CM:
            problems.append({"where": f"paragraph {i}", "param": "first_line_indent_cm",
                             "got": round(pf.first_line_indent.cm, 3), "want": want_ind})
        if want_sp and pf.line_spacing and isinstance(pf.line_spacing, float) \
                and abs(pf.line_spacing - want_sp) > 0.01:
            problems.append({"where": f"paragraph {i}", "param": "line_spacing",
                             "got": pf.line_spacing, "want": want_sp})
    return problems


def repair(docx_path: str | Path, profile: dict, out_path: str | Path) -> dict:
    """Починить документ в профиль -> новый файл. before/after счётчики."""
    d = Document(str(docx_path))
    before = len(validate(docx_path, profile))
    s = d.sections[0]
    m = profile["margins"]
    s.left_margin = Cm(m["left"])
    s.right_margin = Cm(m["right"])
    s.top_margin = Cm(m["top"])
    s.bottom_margin = Cm(m["bottom"])
    pg = profile.get("page", {})
    if pg.get("w_cm"):
        s.page_width = Cm(pg["w_cm"])
    if pg.get("h_cm"):
        s.page_height = Cm(pg["h_cm"])
    fn = d.styles["Normal"].font
    prof_name = profile["font"].get("name")
    prof_size = profile["font"].get("size")
    if prof_name:
        fn.name = prof_name
    if prof_size:
        fn.size = Pt(prof_size)
    if prof_name:
        for where, p in _labeled_paragraphs(d):
            for run in p.runs:
                run.font.name = prof_name
    body = profile.get("body", {})
    want_ind = body.get("first_line_indent_cm")
    want_sp = body.get("line_spacing")
    for p in d.paragraphs:
        pf = p.paragraph_format
        if not p.text.strip():
            continue
        runs_bold = any(r.bold for r in p.runs if r.bold)
        if runs_bold or p.alignment == WD_ALIGN_PARAGRAPH.CENTER or \
                (pf.left_indent and pf.left_indent > 0):
            continue
        if want_ind:
            pf.first_line_indent = Cm(want_ind)
        if want_sp:
            pf.line_spacing = want_sp
    d.save(str(out_path))
    after = len(validate(out_path, profile))
    return {"before": before, "after": after,
            "fixed": max(before - after, 0), "output": str(out_path)}
